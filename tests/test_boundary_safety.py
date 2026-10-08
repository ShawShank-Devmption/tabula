"""Regressions for refusal, grid bounds, immutable loads, and transactional apply."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
import hashlib
import threading

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils.datetime import CALENDAR_MAC_1904
import pytest

from tabula import xlsx
from tabula.tel import compiler


def make_book(tmp_path, configure=None, name="book.xlsx"):
    path = tmp_path / name
    wb = openpyxl.Workbook()
    wb.active.title = "Lists"
    wb.active["A1"] = 1
    wb.create_sheet("Other")
    if configure:
        configure(wb)
    wb.save(path)
    return path


def script_file(tmp_path, source, name="edit.tel"):
    path = tmp_path / name
    path.write_text(source)
    return path


def check(tmp_path, path, source):
    return compiler.check(path, script_file(tmp_path, source), predict_loss=False)


def error_codes(plan):
    return {d.code for d in plan.errors}


def test_rename_refuses_validation_on_another_sheet(tmp_path):
    def setup(wb):
        dv = DataValidation(type="list", formula1="'Lists'!$A$1:$A$3")
        wb["Other"].add_data_validation(dv)
        dv.add("B1")
    path = make_book(tmp_path, setup)
    before = path.read_bytes()
    plan = compiler.apply(path, script_file(tmp_path, 'rename sheet "Lists" to "Regions"'),
                          in_place=True)
    assert "E-STRUCT" in error_codes(plan)
    assert path.read_bytes() == before


@pytest.mark.parametrize("action", ['in "Lists" {\n insert rows 1\n}',
                                    'rename sheet "Lists" to "Regions"'])
@pytest.mark.parametrize("chart_sheet", ["Lists", "Other"])
def test_address_bearing_chart_refuses_structural_edits_anywhere(tmp_path, action, chart_sheet):
    def setup(wb):
        chart = BarChart()
        chart.add_data(Reference(wb["Lists"], min_col=1, min_row=1, max_row=3))
        wb[chart_sheet].add_chart(chart, "D1")
    path = make_book(tmp_path, setup)
    plan, _ = check(tmp_path, path, action)
    assert "E-STRUCT" in error_codes(plan)
    assert "chart" in plan.errors[0].message


@pytest.mark.parametrize("feature", ["formula", "name"])
def test_rename_refuses_unparsed_formula_or_name(tmp_path, feature):
    def setup(wb):
        if feature == "formula":
            wb["Other"]["B1"] = "=Lists!A1#"
        else:
            wb.defined_names.add(DefinedName("Spill", attr_text="Lists!A1#"))
    path = make_book(tmp_path, setup)
    plan, _ = check(tmp_path, path, 'rename sheet "Lists" to "Regions"')
    assert "E-STRUCT" in error_codes(plan)


@pytest.mark.parametrize("axis,address", [("cols", "XFD1"), ("rows", "A1048576")])
@pytest.mark.parametrize("content", ["value", "style", "dimension"])
def test_insertion_refuses_grid_overflow_including_formatting(tmp_path, axis, address, content):
    def setup(wb):
        ws = wb["Lists"]
        if content == "value":
            ws[address] = "edge"
        elif content == "style":
            ws[address].fill = PatternFill("solid", fgColor="FF0000")
        elif axis == "cols":
            ws.column_dimensions["XFD"].width = 20
        else:
            ws.row_dimensions[1048576].height = 20
    path = make_book(tmp_path, setup)
    before = path.read_bytes()
    plan, _ = check(tmp_path, path, f'in "Lists" {{\n insert {axis} {"A" if axis == "cols" else "1"}\n}}')
    assert "E-STRUCT" in error_codes(plan)
    assert path.read_bytes() == before


def test_overflow_guard_follows_prior_delete_and_new_writes(tmp_path):
    path = make_book(tmp_path, lambda wb: setattr(wb["Lists"]["XFD1"], "value", 9))
    source = 'in "Lists" {\n delete cols XFD\n insert cols A\n set XFD1 2\n insert cols A\n}'
    plan, _ = check(tmp_path, path, source)
    assert "E-STRUCT" in error_codes(plan)
    assert len(plan.sim.structure) == 2


def test_structural_output_matches_simulated_cells_and_styles(tmp_path):
    def setup(wb):
        ws = wb["Lists"]
        ws["B2"] = "=A1+1"
        ws["C3"].fill = PatternFill("solid", fgColor="FF0000")
        ws.row_dimensions[3].height = 30
        ws.column_dimensions["C"].width = 25
    path = make_book(tmp_path, setup)
    source = 'in "Lists" {\n insert rows 2\n insert cols B\n}'
    out = tmp_path / "result.xlsx"
    plan = compiler.apply(path, script_file(tmp_path, source), out=out)
    assert plan.applied
    ws = openpyxl.load_workbook(out)["Lists"]
    assert ws["C3"].value == "=A1+1"
    assert ws["D4"].fill.fgColor.rgb == "00FF0000"
    assert ws.row_dimensions[4].height == 30
    assert ws.column_dimensions["D"].width == 25


def test_load_uses_one_immutable_snapshot_for_formulas_caches_and_digest(tmp_path, monkeypatch):
    path = make_book(tmp_path)
    original = path.read_bytes()
    real_load = openpyxl.load_workbook
    calls = 0
    def concurrent_save(*args, **kwargs):
        nonlocal calls
        loaded = real_load(*args, **kwargs)
        calls += 1
        if calls == 1:
            replacement = openpyxl.Workbook()
            replacement.active.title = "External"
            replacement.save(path)
        return loaded
    monkeypatch.setattr(openpyxl, "load_workbook", concurrent_save)
    loaded = xlsx.load(path)
    assert loaded.sha256 == hashlib.sha256(original).hexdigest()
    assert loaded.model.sheet("Lists").cells[(1, 1)].value == 1


@pytest.mark.parametrize("change", ["source", "source_lock", "destination", "destination_lock"])
def test_commit_revalidates_external_changes_after_verification(tmp_path, monkeypatch, change):
    path = make_book(tmp_path)
    output = path if change.startswith("source") else make_book(tmp_path, name="output.xlsx")
    real_verify = xlsx._verify
    external_bytes = None
    def concurrent_save(*args, **kwargs):
        nonlocal external_bytes
        real_verify(*args, **kwargs)
        if change.endswith("lock"):
            output.with_name("~$" + output.name).write_text("Excel")
        else:
            wb = openpyxl.load_workbook(output)
            wb["Lists"]["A1"] = 99
            wb.save(output)
            external_bytes = output.read_bytes()
    monkeypatch.setattr(xlsx, "_verify", concurrent_save)
    plan = compiler.apply(path, script_file(tmp_path, 'set Lists!A1 2'), out=output,
                          in_place=output == path)
    assert not plan.applied
    assert ("E-LOCKED" if change.endswith("lock") else "E-CONFLICT") in error_codes(plan)
    if external_bytes:
        assert output.read_bytes() == external_bytes


def test_cooperating_apply_writers_serialize_before_loading(tmp_path, monkeypatch):
    path = make_book(tmp_path)
    one = script_file(tmp_path, 'set Lists!B1 10', "one.tel")
    two = script_file(tmp_path, 'set Lists!C1 20', "two.tel")
    first_verified = threading.Event()
    release_first = threading.Event()
    real_verify = xlsx._verify
    count_lock = threading.Lock()
    calls = 0
    def pause_first(*args, **kwargs):
        nonlocal calls
        real_verify(*args, **kwargs)
        with count_lock:
            calls += 1
            first = calls == 1
        if first:
            first_verified.set()
            assert release_first.wait(5)
    monkeypatch.setattr(xlsx, "_verify", pause_first)
    with ThreadPoolExecutor(max_workers=2) as pool:
        future_one = pool.submit(compiler.apply, path, one, in_place=True)
        assert first_verified.wait(5)
        future_two = pool.submit(compiler.apply, path, two, in_place=True)
        try:
            # The second writer cannot load the old snapshot and finish while first holds lock.
            assert not future_two.done()
        finally:
            release_first.set()
        assert future_one.result(timeout=5).applied
        assert future_two.result(timeout=5).applied
    ws = openpyxl.load_workbook(path)["Lists"]
    assert (ws["B1"].value, ws["C1"].value) == (10, 20)


def test_1904_epoch_dates_use_workbook_serial_numbers(tmp_path):
    def setup(wb):
        wb.epoch = CALENDAR_MAC_1904
        wb["Lists"]["A1"] = datetime(1904, 1, 2)
    path = make_book(tmp_path, setup)
    plan, loaded = check(tmp_path, path, 'expect Lists!A1 = 1')
    assert loaded.model.sheet("Lists").cells[(1, 1)].value == 1
    assert plan.ok


@pytest.mark.parametrize("action", ['in "Lists" {\n insert rows 1\n}',
                                    'rename sheet "Lists" to "Regions"'])
def test_pivot_elsewhere_refuses_structural_edits(tmp_path, action):
    from openpyxl.pivot.table import TableDefinition, Location
    from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource
    def setup(wb):
        pivot = TableDefinition(name="Pivot", cacheId=1, dataCaption="Values",
                                location=Location(ref="A3:A4", firstHeaderRow=1,
                                                  firstDataRow=1, firstDataCol=1))
        pivot.cache = CacheDefinition(cacheSource=CacheSource(
            type="worksheet", worksheetSource=WorksheetSource(ref="A1:A3", sheet="Lists")))
        wb["Other"].add_pivot(pivot)
    path = make_book(tmp_path, setup)
    plan, _ = check(tmp_path, path, action)
    assert "E-STRUCT" in error_codes(plan)
    assert "pivot" in plan.errors[0].message


def test_synced_count_includes_only_written_cells_not_untouched_moved_cells(tmp_path):
    path = make_book(tmp_path)
    plan = compiler.apply(path, script_file(tmp_path, 'in "Lists" {\n insert rows 1\n set B1 2\n}'),
                          out=tmp_path / "output.xlsx")
    assert plan.applied
    assert plan.synced == 1


def test_postwrite_checks_untouched_structurally_moved_values(tmp_path, monkeypatch):
    path = make_book(tmp_path)
    before = path.read_bytes()
    real_save = openpyxl.Workbook.save
    def corrupt_moved_cell(wb, destination):
        real_save(wb, destination)
        saved = openpyxl.load_workbook(destination)
        saved["Lists"]["A2"] = 99
        real_save(saved, destination)
    monkeypatch.setattr(openpyxl.Workbook, "save", corrupt_moved_cell)
    plan = compiler.apply(path, script_file(tmp_path, 'in "Lists" {\n insert rows 1\n}'),
                          in_place=True)
    assert "E-VERIFY" in error_codes(plan)
    assert path.read_bytes() == before


@pytest.mark.parametrize("feature", ["formula", "name"])
def test_indirect_text_addresses_block_rename_and_structural_edits(tmp_path, feature):
    def setup(wb):
        if feature == "formula":
            wb["Other"]["B1"] = '=INDIRECT("Lists!A1")'
        else:
            wb.defined_names.add(DefinedName("Dynamic", attr_text='INDIRECT("Lists!A1")'))
    path = make_book(tmp_path, setup)
    plan, _ = check(tmp_path, path, 'rename sheet "Lists" to "Regions"')
    assert "E-STRUCT" in error_codes(plan)


def test_unparsed_formula_is_listed_as_unverified_after_edit(tmp_path):
    path = make_book(tmp_path, lambda wb: setattr(wb["Other"]["B1"], "value", "=Lists!A1#"))
    plan, _ = check(tmp_path, path, 'set Lists!A1 2')
    assert "Other!B1" in plan.sim.unverified


@pytest.mark.parametrize("target", ["source", "existing_output", "new_output"])
def test_external_save_during_simulation_is_preserved(tmp_path, monkeypatch, target):
    path = make_book(tmp_path)
    output = path if target == "source" else tmp_path / "output.xlsx"
    if target == "existing_output":
        make_book(tmp_path, name="output.xlsx")
    real_check = compiler.check
    external_bytes = None
    def concurrent_save(*args, **kwargs):
        nonlocal external_bytes
        result = real_check(*args, **kwargs)
        wb = openpyxl.Workbook()
        wb.active.title = "External"
        wb.active["A1"] = "keep this external edit"
        wb.save(output)
        external_bytes = output.read_bytes()
        return result
    monkeypatch.setattr(compiler, "check", concurrent_save)
    plan = compiler.apply(path, script_file(tmp_path, 'set Lists!A1 2'), out=output,
                          in_place=target == "source")
    assert "E-CONFLICT" in error_codes(plan)
    assert output.read_bytes() == external_bytes


def test_postwrite_checks_unexpected_structural_cells(tmp_path, monkeypatch):
    path = make_book(tmp_path)
    before = path.read_bytes()
    real_save = openpyxl.Workbook.save
    def corrupt_extra_cell(wb, destination):
        real_save(wb, destination)
        saved = openpyxl.load_workbook(destination)
        saved["Lists"]["XFE1"] = 99
        real_save(saved, destination)
    monkeypatch.setattr(openpyxl.Workbook, "save", corrupt_extra_cell)
    plan = compiler.apply(path, script_file(tmp_path, 'in "Lists" {\n insert rows 1\n}'),
                          in_place=True)
    assert "E-VERIFY" in error_codes(plan)
    assert path.read_bytes() == before


def test_refused_rename_stops_ops_bound_to_new_sheet_name(tmp_path):
    def setup(wb):
        dv = DataValidation(type="list", formula1="'Lists'!$A$1:$A$3")
        wb["Other"].add_data_validation(dv)
        dv.add("B1")
    path = make_book(tmp_path, setup)
    before = path.read_bytes()
    plan = compiler.apply(path, script_file(tmp_path, 'rename sheet "Lists" to "Regions"\nset Regions!A1 2'),
                          in_place=True)
    assert "E-STRUCT" in error_codes(plan)
    assert not plan.applied
    assert path.read_bytes() == before
