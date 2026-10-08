"""TEL compiler tests: parsing, semantic checks, IR optimisation, simulation, apply."""
import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "demo"))

import build_workbook  # noqa: E402

from tabula.tel import compiler  # noqa: E402
from tabula.tel.ir import ClearCell, Expect, SetCell, eliminate_dead_writes  # noqa: E402
from tabula.tel.parser import parse_script  # noqa: E402
from tabula.tel.plan import render_json  # noqa: E402
from tabula import xlsx  # noqa: E402


@pytest.fixture()
def book(tmp_path) -> Path:
    path = tmp_path / "sales.xlsx"
    build_workbook.build(path)
    return path


def run(book, source, tmp_path, **kw):
    script = tmp_path / "edit.tel"
    script.write_text(source)
    plan, loaded = compiler.check(book, script, **kw)
    return plan, loaded


def codes(plan):
    return [d.code for d in plan.sorted_diagnostics()]


# ---------------------------------------------------------------- parser
def test_parser_recovers_line_by_line():
    s = parse_script('set A1 5\nset A1 [1,\nbogus x\nin "S" {\n  clear A1\n')
    assert [d.line for d in s.diagnostics] == [2, 3, 4]   # bad list, bad keyword, unclosed block
    assert len(s.statements) == 2                          # set + the (unclosed) in-block


def test_formula_error_positions_map_into_the_script():
    s = parse_script("set A1 =SUM(A2:A3")
    d = s.diagnostics[0]
    assert (d.code, d.line, d.col) == ("E-SYNTAX", 1, 18)


# ---------------------------------------------------------------- semantic checks
@pytest.mark.parametrize("source,code", [
    ("set Summry!B8 1", "E-SHEET"),
    ("set B3 5", "E-NOSHEET"),
    ('in "Sales" {\n set G2 =E2*TaxRat\n}', "E-NAME"),
    ('in "Sales" {\n set G2 =SUMM(E2:E3)\n}', "E-FUNC"),
    ('in "Sales" {\n set G2 =ROUND(E2)\n}', "E-ARITY"),
    ('in "Sales" {\n set G2 =E2:E21*2\n}', "E-TYPE"),
    ('in "Sales" {\n set A1:C1 ["x","y"]\n}', "E-SHAPE"),
    ('in "Sales" {\n let t = A1\n let t = A2\n}', "E-DECL"),
    ('in "Sales" {\n set A:A 0\n}', "E-SIZE"),
    ('in "Sales" {\n set H1 =A0+1\n}', "E-REF"),          # row 0 does not exist
    ('in "Sales" {\n set H1 =SUM(A1\n}', "E-SYNTAX"),
])
def test_front_end_errors(book, tmp_path, source, code):
    plan, _ = run(book, source, tmp_path)
    assert code in codes(plan) and not plan.ok


def test_suggestions_for_typos(book, tmp_path):
    plan, _ = run(book, "set Summry!B8 1", tmp_path)
    assert plan.errors[0].hint == "did you mean 'Summary'?"


def test_let_binding_is_inlined_and_scoped(book, tmp_path):
    source = ('let rate = 0.2\nin "Sales" {\n  let rev = E2:E21\n  set G22 =SUM(rev)*rate\n}\n'
              'expect Sales!G22 = Sales!E22*0.2')
    plan, _ = run(book, source, tmp_path)
    assert plan.ok, codes(plan)
    assert plan.sim.writes[0].after_raw == "=SUM($E$2:$E$21)*0.2"   # let inlines as absolute
    assert plan.sim.expects[0]["status"] == "pass"


def test_scope_warning_for_unqualified_refs_on_another_sheet(book, tmp_path):
    plan, _ = run(book, 'in "Sales" {\n  set Summary!B9 =E22\n}', tmp_path)
    assert "W-SCOPE" in codes(plan)


# ---------------------------------------------------------------- IR optimisation
def test_dead_write_elimination_respects_barriers():
    ops = [SetCell("S", 1, 1, value=1, line=1), SetCell("S", 1, 1, value=2, line=2),
           Expect("S", None, "", line=3), SetCell("S", 1, 1, value=3, line=4),
           ClearCell("S", 1, 1, line=5)]
    live, dead = eliminate_dead_writes(ops)
    assert [op.line for op in live] == [2, 3, 5]
    assert dead == {(1, 2): 1, (4, 5): 1}


# ---------------------------------------------------------------- simulation
def test_fill_and_predicted_values(book, tmp_path):
    plan, _ = run(book, 'in "Sales" {\n  set G2:G21 =E2*TaxRate\n}', tmp_path)
    assert plan.ok
    g3 = next(w for w in plan.sim.writes if w.address == "Sales!G3")
    assert g3.after_raw == "=E3*TaxRate" and g3.after_value == pytest.approx(18 * 250 * 0.18)


def test_cycle_is_rejected_with_path(book, tmp_path):
    plan, _ = run(book, 'in "Sales" {\n  set E22 =SUM(E2:E22)\n}', tmp_path)
    err = next(d for d in plan.errors if d.code == "E-CYCLE")
    assert "Sales!E22 -> Sales!E22" in err.message and err.line == 2


def test_hard_coding_and_impact(book, tmp_path):
    plan, _ = run(book, 'in "Sales" {\n  set E2 1000\n}', tmp_path)
    assert "W-OVERWRITE-FORMULA" in codes(plan) and plan.ok
    affected = {c.address for c in plan.sim.affected}
    assert {"Sales!F2", "Sales!E22", "Summary!B3"} <= affected


def test_failed_expect_explains_both_sides(book, tmp_path):
    plan, _ = run(book, "expect Summary!B3 = 1", tmp_path)
    err = plan.errors[0]
    assert err.code == "E-EXPECT" and "right side 1" in err.message


def test_insert_rows_relocates_everything(book, tmp_path):
    src = (ROOT / "demo" / "reference" / "t2_insert_row.tel").read_text()
    plan, loaded = run(book, src, tmp_path)
    assert plan.ok, [d.render() for d in plan.diagnostics]
    wb = loaded.model
    assert wb.sheet("Summary").cells[(2, 3)].raw == "=Sales!E23"
    assert wb.names[(None, "SALESREVENUE")].text == "Sales!$E$2:$E$22"
    assert all(e["status"] == "pass" for e in plan.sim.expects)


def test_unsimulated_function_is_flagged_not_guessed(book, tmp_path):
    plan, _ = run(book, 'in "Summary" {\n  set B9 =VLOOKUP("Asha Rao",Sales!B2:E21,4,FALSE)\n}',
                  tmp_path)
    assert "W-UNSIMULATED" in codes(plan) and plan.ok
    assert plan.sim.writes[0].unverified


def test_allowlist(book, tmp_path):
    plan, _ = run(book, 'in "Sales" {\n  set G2 1\n}', tmp_path, allow_specs=["Summary"])
    assert "E-PERM" in codes(plan)


def test_json_plan_is_versioned_and_bounded(book, tmp_path):
    plan, _ = run(book, 'in "Sales" {\n  set G2:G300 =E2*2\n}', tmp_path)
    out = json.loads(render_json(plan))
    assert out["schema"] == 1 and out["writes"]["count"] == 299
    assert len(out["writes"]["items"]) == 50


# ---------------------------------------------------------------- apply / I/O
def test_apply_writes_verified_file(book, tmp_path):
    script = tmp_path / "t.tel"
    script.write_text('rename sheet "Sales" to "Q3 Sales"')
    out = tmp_path / "out.xlsx"
    plan = compiler.apply(book, script, out=out)
    assert plan.applied, [d.render() for d in plan.diagnostics]
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Inputs", "Q3 Sales", "Summary"]
    assert wb["Summary"]["B3"].value == "='Q3 Sales'!E22"
    assert wb.defined_names["SalesRevenue"].attr_text == "'Q3 Sales'!$E$2:$E$21"


def test_check_never_writes_and_errors_block_apply(book, tmp_path):
    before = xlsx.sha256(book)
    script = tmp_path / "bad.tel"
    script.write_text("set Nope!A1 1")
    plan = compiler.apply(book, script, in_place=True)
    assert not plan.applied and xlsx.sha256(book) == before


def test_conflict_and_lock_guards(book, tmp_path):
    script = tmp_path / "t.tel"
    script.write_text('in "Sales" {\n  set G1 "Tax"\n}')
    plan = compiler.apply(book, script, in_place=True, if_unchanged="0" * 64)
    assert "E-CONFLICT" in codes(plan) and not plan.applied
    (book.parent / f"~${book.name}").write_text("lock")
    plan = compiler.apply(book, script, in_place=True)
    assert "E-LOCKED" in codes(plan) and not plan.applied


def test_failed_save_leaves_input_untouched(book, tmp_path, monkeypatch):
    before = xlsx.sha256(book)
    script = tmp_path / "t.tel"
    script.write_text('in "Sales" {\n  set G1 "Tax"\n}')

    def boom(*a, **k):
        raise OSError("disk full")
    monkeypatch.setattr(xlsx, "_verify", boom)
    plan = compiler.apply(book, script, in_place=True)
    assert not plan.applied and "E-WRITE" in codes(plan)
    assert xlsx.sha256(book) == before
    assert not list(book.parent.glob(".tabula-*"))       # temp file cleaned up


def test_literal_text_starting_with_equals_stays_text(book, tmp_path):
    script = tmp_path / "t.tel"
    script.write_text('in "Sales" {\n  set H1 "=not a formula"\n}')
    out = tmp_path / "out.xlsx"
    assert compiler.apply(book, script, out=out).applied
    cell = openpyxl.load_workbook(out)["Sales"]["H1"]
    assert cell.value == "=not a formula" and cell.data_type == "s"


def _inject_part(book: Path, name: str, data: bytes) -> None:
    """Add a zip part openpyxl does not model (it will be dropped on save)."""
    import zipfile
    tmp = book.with_suffix(".tmp")
    with zipfile.ZipFile(book) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr(name, data)
    tmp.replace(book)


def test_fidelity_guard_refuses_lossy_save(book, tmp_path):
    _inject_part(book, "customXml/item1.xml", b"<root/>")
    before = xlsx.sha256(book)
    script = tmp_path / "t.tel"
    script.write_text('in "Sales" {\n  set G1 "Tax"\n}')
    plan, _ = compiler.check(book, script)
    assert "W-LOSSY" in codes(plan) and plan.ok            # check warns early
    plan = compiler.apply(book, script, in_place=True)
    assert "E-FIDELITY" in codes(plan) and not plan.applied
    assert xlsx.sha256(book) == before                      # input untouched
    plan = compiler.apply(book, script, out=tmp_path / "lossy.xlsx", allow_lossy=True)
    assert plan.applied
