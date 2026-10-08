"""Evaluation regressions: checker behavior must reject realistic corruptions."""
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'demo'))


def test_suite_has_fifteen_frozen_holdout_tasks():
    import suite
    assert len(suite.TASKS) == 15
    assert len({t.workbook for t in suite.TASKS}) >= 5
    assert not any(t.id.startswith('t') for t in suite.TASKS)
    assert len({t.id for t in suite.TASKS}) == 15


def check(task, path):
    script = "import sys,json;sys.path.insert(0,'demo');from suite_evaluate import check_task; print(json.dumps(check_task(sys.argv[1],sys.argv[2])))"
    result = subprocess.run([str(ROOT / 'demo/.venv/bin/python'), '-c', script, task.id, str(path)], cwd=ROOT, capture_output=True, text=True, check=True)
    return json.loads(result.stdout)


def test_reference_solutions_and_mutations(tmp_path):
    import suite
    for task in suite.TASKS:
        path = tmp_path / f'{task.id}.xlsx'
        suite.build(task, path)
        suite.solve(task, path)
        checks = check(task, path)
        assert checks and all(x['ok'] for x in checks), (task.id, checks)
        wb = openpyxl.load_workbook(path)
        wb['Directory']['C2'] = 'CORRUPTED unrelated annotation'
        wb.save(path)
        assert any(not x['ok'] and 'preservation' in x['check'] for x in check(task, path))


def test_constant_formula_and_style_mutations_fail(tmp_path):
    import suite
    task = suite.TASKS[0]
    path = tmp_path / 'book.xlsx'
    suite.build(task, path)
    suite.solve(task, path)
    wb = openpyxl.load_workbook(path)
    spec = suite.SPECS[task.workbook]
    addr = suite.positions(task)['amount_cells'][0]
    wb[spec.sheet][addr] = '=17*11'
    wb.save(path)
    assert any(not c['ok'] and 'liveness' in c['check'] for c in check(task, path))
    suite.build(task, path)
    suite.solve(task, path)
    wb = openpyxl.load_workbook(path)
    wb['Directory']['C2'].number_format = '0.000'
    wb.save(path)
    assert any(not c['ok'] and 'preservation' in c['check'] for c in check(task, path))


def test_exit_zero_is_not_a_success_claim():
    from suite_evaluate import classify
    checks = [{'check':'preservation', 'ok':False, 'detail':'changed record'}]
    assert classify({'returncode':0, 'final_message':'All done.'}, checks)['outcome'] == 'unverified'
    assert not classify({'returncode':0, 'final_message':'All done.'}, checks)['silent_corruption']
    assert classify({'returncode':0, 'final_message':'EVAL_STATUS: SUCCESS'}, checks)['silent_corruption']
    assert classify({'returncode':0, 'final_message':'EVAL_STATUS: REFUSAL'}, checks)['outcome'] == 'failure'
    assert classify({'returncode':0, 'final_message':'EVAL_STATUS: REFUSAL'}, [{'ok':True}])['outcome'] == 'safe_refusal'


@pytest.mark.parametrize('args', [['--trials','0'], ['--arms','typo'], ['--tasks','typo'], ['--timeout','0'], ['--budget','nan'], ['--parallel','0']])
def test_runner_rejects_invalid_arguments_before_creating_runs(tmp_path, args):
    target = tmp_path / 'runs'
    result = subprocess.run([sys.executable, str(ROOT/'demo/suite_run.py'), '--driver','reference', '--out',str(target), *args], capture_output=True, text=True)
    assert result.returncode != 0
    assert not target.exists()


def test_old_tax_demo_unrelated_salesperson_corruption_rejected(tmp_path):
    path = tmp_path / 'book.xlsx'
    import shutil
    shutil.copy(ROOT/'demo/runs/reference/t1_tax_column__default__1/book.xlsx', path)
    wb = openpyxl.load_workbook(path)
    wb['Sales']['B2'] = 'CORRUPTED SALESPERSON'
    wb.save(path)
    script = "import sys,json;sys.path.insert(0,'demo');from evaluate import check_task;print(json.dumps(check_task('t1_tax_column',__import__('pathlib').Path(sys.argv[1]))))"
    result = subprocess.run([str(ROOT/'demo/.venv/bin/python'), '-c', script, str(path)], cwd=ROOT, capture_output=True, text=True, check=True)
    checks = json.loads(result.stdout)
    assert any(not c['ok'] and 'data rows' in c['check'] for c in checks)
