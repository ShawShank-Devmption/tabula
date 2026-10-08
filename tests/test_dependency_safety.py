"""Safety of predicted values: hidden reads, cached unknowns and edit invalidation."""
from pathlib import Path
import random
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from tabula import xlsx
from tabula.engine import Sheet
from tabula.tel.compiler import check


def cached_book(tmp_path, formula, names=()):
    """Synthetic caches with hand-known values, not an Excel differential corpus."""
    w = openpyxl.Workbook()
    s = w.active
    s.title = 'S'
    s['A1'], s['B1'], s['C1'] = 10, formula, '=B1*2'
    for name, value in names:
        w.defined_names.add(DefinedName(name, attr_text=value))
    path = tmp_path / 'book.xlsx'
    w.save(path)
    with ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    tree = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    for cell in tree.findall('.//m:c', ns):
        if cell.attrib['r'] in ('B1', 'C1'):
            cell.find('m:v', ns).text = {'B1': '10', 'C1': '20'}[cell.attrib['r']]
    parts['xl/worksheets/sheet1.xml'] = ET.tostring(tree)
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    return path


@pytest.mark.parametrize('formula,names', [
    ('=INDIRECT("A1")', ()),
    ('=SUM(A1#)', ()),
    ('=Hidden', (('Hidden', 'INDIRECT("A1")'),)),
    ('=Hidden', (('Hidden', 'SUM(S!A1#)'),)),
])
def test_edit_invalidates_unknown_dependencies_and_expectations(tmp_path, formula, names):
    path = cached_book(tmp_path, formula, names)
    plan, loaded = check(path, tmp_path / 'edit.tel',
                         source='set A1 99\nexpect C1 = 20', predict_loss=False)
    assert plan.sim.expects[0]['status'] == 'unverified'
    assert {'S!B1', 'S!C1'} <= set(plan.sim.unverified)
    assert loaded.model.sheet('S').cells[(3, 1)].unverified


def test_empty_check_keeps_original_cache_verified(tmp_path):
    path = cached_book(tmp_path, '=INDIRECT("A1")')
    plan, _ = check(path, tmp_path / 'edit.tel', source='expect C1 = 20', predict_loss=False)
    assert plan.sim.expects[0]['status'] == 'pass'


def test_uncertainty_is_propagated_even_when_error_handler_masks_value(tmp_path):
    path = cached_book(tmp_path, '=INDIRECT("A1")')
    plan, _ = check(path, tmp_path / 'edit.tel',
                    source='set A1 99\nset C1 =IFERROR(B1*2,0)\nexpect C1=20', predict_loss=False)
    assert plan.sim.expects[0]['status'] == 'unverified'


def test_sumif_effective_range_invalidates_dependent(tmp_path):
    w = openpyxl.Workbook()
    s = w.active
    s.title = 'S'
    for r in range(1, 4):
        s.cell(r, 1, 1)
        s.cell(r, 2, r * 10)
    s['C1'] = '=SUMIF(A1:A3,">0",B1)'
    path = tmp_path / 'book.xlsx'
    w.save(path)
    plan, loaded = check(path, tmp_path / 'edit.tel',
                         source='set B3 100\nexpect C1=130', predict_loss=False)
    assert plan.ok
    assert loaded.model.sheet('S').cells[(3, 1)].value == 130
    assert 'S!C1' in {c.address for c in plan.sim.affected}


def test_sumif_implicit_self_reference_is_a_cycle():
    s = Sheet()
    s.set('A1', '1')
    s.set('A2', '1')
    s.set('B1', '3')
    s.set('B2', '=SUMIF(A1:A2,">0",B1)')
    assert s.engine.cyclic, 'B2 is inside the effective sum range B1:B2'


def test_named_cross_sheet_sumif_effective_range(tmp_path):
    w = openpyxl.Workbook()
    a = w.active
    a.title = 'Criteria'
    b = w.create_sheet('Amounts')
    for r in range(1, 4):
        a.cell(r, 1, 1)
        b.cell(r, 2, r * 10)
    w.defined_names.add(DefinedName('Rows', attr_text='Criteria!$A$1:$A$3'))
    w.defined_names.add(DefinedName('Start', attr_text='Amounts!$B$1'))
    a['D1'] = '=SUMIF(Rows,">0",Start)'
    path = tmp_path / 'named.xlsx'
    w.save(path)
    plan, _ = check(path, tmp_path / 'e.tel', source='set Amounts!B3 100\nexpect Criteria!D1=130',
                    predict_loss=False)
    assert plan.ok


def test_random_edits_incremental_equal_fresh_full_recompute():
    # The full path is an independent schedule, and literal arithmetic is checked too.
    rng = random.Random(20261008)
    s = Sheet()
    for r in range(1, 9):
        s.set(f'A{r}', '1')
        s.set(f'B{r}', str(r))
    s.set('D1', '=SUMIF(A1:A8,">0",B1)')
    s.set('E1', '=D1*2')
    values = list(range(1, 9))
    for _ in range(30):
        row, value = rng.randint(1, 8), rng.randint(-100, 100)
        values[row - 1] = value
        s.set(f'B{row}', str(value))
        incremental = (s.value('D1'), s.value('E1'))
        assert incremental == (sum(values), 2 * sum(values))
        s.engine.recompute()
        assert incremental == (s.value('D1'), s.value('E1'))


def test_expect_cannot_mask_an_unsimulated_function_as_a_known_value(tmp_path):
    path = cached_book(tmp_path, '=INDIRECT("A1")')
    plan, _ = check(path, tmp_path / 'edit.tel',
                    source='expect IFERROR(INDIRECT("A1"),10)=10', predict_loss=False)
    assert plan.sim.expects[0]['status'] == 'unverified'


def test_cyclic_defined_names_do_not_crash_load(tmp_path):
    path = cached_book(tmp_path, '=First', [('First', 'Second'), ('Second', 'First')])
    plan, _ = check(path, tmp_path / 'edit.tel',
                    source='set A1 11\nexpect C1=20', predict_loss=False)
    assert plan.sim.expects[0]['status'] == 'unverified'
