"""Regression tests for bugs found in the Revision 2 code review (one test per finding)."""
from pathlib import Path

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from tabula import xlsx
from tabula.engine import Sheet
from tabula.tel import compiler
from tabula.values import display


def make(tmp_path: Path, sheets: dict, names: dict | None = None, name="book") -> Path:
    wb = openpyxl.Workbook()
    for i, (title, cells) in enumerate(sheets.items()):
        ws = wb.active if i == 0 else wb.create_sheet(title)
        ws.title = title
        for coord, value in cells.items():
            ws[coord] = value
    for n, text in (names or {}).items():
        wb.defined_names[n] = DefinedName(n, attr_text=text)
    path = tmp_path / f"{name}.xlsx"
    wb.save(path)
    return path


def check(tmp_path, book, source):
    script = tmp_path / "edit.tel"
    script.write_text(source)
    return compiler.check(book, script, predict_loss=False)


def apply(tmp_path, book, source, out="out.xlsx"):
    script = tmp_path / "edit.tel"
    script.write_text(source)
    return compiler.apply(book, script, out=tmp_path / out)


def codes(plan):
    return [d.code for d in plan.diagnostics]


# 1. values through whole-column ranges / relocated names must be recomputed after delete
def test_structural_edit_recomputes_names_and_whole_columns(tmp_path):
    book = make(tmp_path, {"Data": {"A1": 10, "A2": 20, "A3": 30},
                           "Summary": {"A1": "=SUM(Total)", "A2": "=SUM(Data!A:A)",
                                       "A3": "=COUNTA(Data!1:1)"}},
                {"Total": "Data!$A$1:$A$3"})
    plan, _ = check(tmp_path, book, 'in "Data" {\n  delete rows 2\n}\n'
                                    "expect Summary!A1 = 40\nexpect Summary!A2 = 40")
    assert plan.ok, [d.render() for d in plan.diagnostics]
    assert {c.address for c in plan.sim.affected} >= {"Summary!A1", "Summary!A2"}


# 2. a let-bound address must not drift when a formula is filled down
def test_let_names_do_not_shift_under_fill(tmp_path):
    book = make(tmp_path, {"S": {"A1": 1, "A2": 2, "A3": 3}, "Cfg": {"B1": 10}})
    plan, _ = check(tmp_path, book, 'let rate = Cfg!B1\nlet amounts = S!$A$1:$A$3\n'
                                    'in "S" {\n  set C1:C3 =A1*rate\n  set D1:D3 =SUM(amounts)\n}')
    assert plan.ok
    raws = {w.address: w.after_raw for w in plan.sim.writes}
    assert raws["S!C3"] == "=A3*Cfg!$B$1"
    assert raws["S!D3"] == "=SUM($A$1:$A$3)"


# 3. a defined name used as a target follows earlier structural edits
def test_defined_name_target_follows_insert(tmp_path):
    book = make(tmp_path, {"Data": {"A1": "Rate", "A2": 0.1}}, {"Rate": "Data!$A$2"})
    plan = apply(tmp_path, book, 'in "Data" {\n  insert rows 1\n  set Rate 0.2\n}')
    assert plan.applied
    ws = openpyxl.load_workbook(tmp_path / "out.xlsx")["Data"]
    assert (ws["A2"].value, ws["A3"].value) == ("Rate", 0.2)


# 4. a pre-existing cycle must not crash or be blamed on moved cells
@pytest.mark.parametrize("cells,source", [
    ({"C3": "=C3+1"}, 'in "S" {\n  delete rows 3\n}'),
    ({"B4": "=1+1", "B5": "=B5+1"}, 'in "S" {\n  insert rows 1\n}'),
])
def test_existing_cycles_survive_structural_edits(tmp_path, cells, source):
    book = make(tmp_path, {"S": cells})
    plan, _ = check(tmp_path, book, source)
    assert "E-CYCLE" not in codes(plan) and plan.ok


# 5 + 6. single-cell references and blanks in aggregates/criteria, as Excel
@pytest.mark.parametrize("formula,expected", [
    ("=AVERAGE(A1,A2)", "10"), ("=MIN(A1,A2)", "10"), ("=SUM(A1,A3)", "10"),
    ("=MAX(A1,A3)", "10"), ("=COUNT(A1,A4)", "1"), ('=COUNTIF(A3,"abc")', "1"),
    ('=SUMIF(A1:A3,"abc",B1)', "5"), ('=COUNTIF(A1:A5,"")', "2"),
    ('=COUNTIF(A1:A5,"<>abc")', "4"), ('=SUMIF(D1:D3,"<>x",E1:E3)', "6"),
])
def test_references_and_blanks_follow_excel(formula, expected):
    s = Sheet()
    for ref, raw in {"A1": "10", "A3": "abc", "A4": "TRUE", "B3": "5", "D1": "x",
                     "E1": "1", "E2": "2", "E3": "4"}.items():
        s.set(ref, raw)
    s.set("Z1", formula)
    assert display(s.value("Z1")) == expected


# 7. renaming a sheet to a different case
def test_rename_changing_only_case(tmp_path):
    book = make(tmp_path, {"Sales": {"A1": 1}, "Summary": {"A1": "=Sales!A1"}})
    plan = apply(tmp_path, book, 'rename sheet "Sales" to "SALES"')
    assert plan.applied, [d.message for d in plan.diagnostics]
    wb = openpyxl.load_workbook(tmp_path / "out.xlsx")
    assert wb.sheetnames == ["SALES", "Summary"] and wb["Summary"]["A1"].value == "=SALES!A1"


# 8. text that looks like an error code stays text (both directions)
def test_error_like_text_stays_text(tmp_path):
    book = make(tmp_path, {"S": {"A2": "=ISERROR(A1)"}})
    plan = apply(tmp_path, book, 'in "S" {\n  set A1 "#N/A"\n}')
    assert plan.applied
    cell = openpyxl.load_workbook(tmp_path / "out.xlsx")["S"]["A1"]
    assert (cell.value, cell.data_type) == ("#N/A", "s")
    loaded = xlsx.load(tmp_path / "out.xlsx")
    assert loaded.model.sheet("S").cells[(1, 2)].value is False


# 9. overflow is #NUM!, not a crash
def test_overflow_is_num_error(tmp_path):
    s = Sheet()
    s.set("A1", "=1E300*1E300")
    assert display(s.value("A1")) == "#NUM!"
    book = make(tmp_path, {"S": {"A1": 1}})
    plan, _ = check(tmp_path, book, 'in "S" {\n  set B1 1e400\n}')
    assert "E-SYNTAX" in codes(plan)


# 10. structural edits refused where references outside formulas would go stale
def test_structural_edit_refused_when_validation_elsewhere_points_here(tmp_path):
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    lists = wb.active
    lists.title = "Lists"
    for i, v in enumerate(["Region", "North", "South"], 1):
        lists[f"A{i}"] = v
    form = wb.create_sheet("Form")
    dv = DataValidation(type="list", formula1="=Lists!$A$2:$A$3")
    dv.add("B2")
    form.add_data_validation(dv)
    path = tmp_path / "dv.xlsx"
    wb.save(path)
    plan, _ = check(tmp_path, path, 'in "Lists" {\n  insert rows 1\n}')
    assert "E-STRUCT" in codes(plan)


def test_hidden_rows_move_with_their_data(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 7):
        ws[f"A{r}"] = f"row{r}"
    ws.row_dimensions[5].hidden = True
    path = tmp_path / "hidden.xlsx"
    wb.save(path)
    assert apply(tmp_path, path, 'in "S" {\n  insert rows 2\n}').applied
    out = openpyxl.load_workbook(tmp_path / "out.xlsx")["S"]
    assert out["A6"].value == "row5" and out.row_dimensions[6].hidden
    assert not out.row_dimensions[5].hidden


# 11. functions newer than Excel 2007 are stored with the _xlfn. prefix
def test_new_functions_get_storage_prefix(tmp_path):
    book = make(tmp_path, {"S": {"A1": "a", "B1": "b"}})
    plan = apply(tmp_path, book, 'in "S" {\n  set C1 =CONCAT(A1,B1)\n}')
    assert plan.applied
    assert openpyxl.load_workbook(tmp_path / "out.xlsx")["S"]["C1"].value == "=_xlfn.CONCAT(A1,B1)"
    loaded = xlsx.load(tmp_path / "out.xlsx")            # and it reads back as CONCAT
    assert loaded.model.sheet("S").cells[(3, 1)].value == "ab"


# minor findings
def test_empty_string_literal_clears(tmp_path):
    book = make(tmp_path, {"S": {"A1": 5}})
    assert apply(tmp_path, book, 'in "S" {\n  set A1 ""\n}').applied


@pytest.mark.parametrize("formula,expected", [("=0^0", "#NUM!"), ("=0.1+0.2=0.3", "TRUE")])
def test_minor_excel_semantics(formula, expected):
    s = Sheet()
    s.set("A1", formula)
    assert display(s.value("A1")) == expected
