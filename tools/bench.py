"""Benchmark: incremental vs full recomputation and end-to-end `check` time (NFR-2).

Builds a workbook with ~10,000 formulas in several topologies, writes it to a
temporary .xlsx, then measures:
  - full recompute of every formula (engine only)
  - a bounded single-cell edit simulated through the TEL compiler (engine only)
  - `check` end to end, including loading the file
Run:  python3 tools/bench.py
"""
from __future__ import annotations

import sys
import tempfile
import time
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tabula import xlsx  # noqa: E402
from tabula.tel import compiler  # noqa: E402
from tabula.tel.simulate import simulate  # noqa: E402

ROWS = 2500


def build(path: Path) -> int:
    """Four columns of formulas over ROWS rows: per-row, running chain, windowed SUM, IF."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, ROWS + 1):
        ws.cell(r, 1, r % 97)                                   # A: input
        ws.cell(r, 2, f"=A{r}*2")                               # B: per-row
        ws.cell(r, 3, "=B1" if r == 1 else f"=C{r - 1}+B{r}")   # C: running chain (depth ROWS)
        lo = max(1, r - 9)
        ws.cell(r, 4, f"=SUM(B{lo}:B{r})")                      # D: windowed range
        ws.cell(r, 5, f'=IF(D{r}>500,"high","low")')            # E: conditional
    ws.cell(ROWS + 1, 3, f"=SUM(C1:C{ROWS})")
    wb.save(path)
    return ROWS * 4 + 1


def main() -> None:
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "bench.xlsx"
        n = build(path)
        t = time.perf_counter()
        loaded = xlsx.load(path)
        load_s = time.perf_counter() - t
        engine = loaded.engine
        stats = engine.recompute()
        print(f"formulas: {n}   load + initial compute: {load_s:.2f} s")
        print(f"full recompute: {stats['formulas_recomputed']} formulas in {stats['recompute_ms']:.1f} ms")

        script = Path(d) / "edit.tel"
        for label, source in [
            ("edit near the bottom (bounded)", f"set Data!A{ROWS - 5} 1000"),
            ("edit at the top (chain to the end)", "set Data!A1 1000"),
        ]:
            loaded = xlsx.load(path)
            ops = compiler.compile_source(source, loaded.model)[1]
            t = time.perf_counter()
            res = simulate(loaded.engine, ops)
            sim_ms = (time.perf_counter() - t) * 1000
            print(f"{label}: recomputed {res.stats['formulas_recomputed']} of "
                  f"{res.stats['total_formulas']} formulas; simulate {sim_ms:.1f} ms "
                  f"(recompute {res.stats['recompute_ms']:.1f} ms)")
        script.write_text(f"set Data!A{ROWS - 5} 1000")
        t = time.perf_counter()
        plan, _ = compiler.check(path, script, predict_loss=False)
        print(f"check end to end (load + compile + simulate): {time.perf_counter() - t:.2f} s, "
              f"ok={plan.ok}")


if __name__ == "__main__":
    main()
