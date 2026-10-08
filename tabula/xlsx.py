"""The .xlsx boundary (design.md section 7) -- the ONLY module that imports openpyxl.

openpyxl reads and writes the file format.  It never evaluates formulas and
never rewrites references; Tabula does both.  Writing is transactional:
replay structure -> sync contents -> save to a temp file -> fidelity guard ->
verify by reloading -> atomic rename.  Any failure leaves the input untouched.
"""
from __future__ import annotations

import hashlib
import os
import re
import tempfile
import warnings
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
from openpyxl.utils import FORMULAE, column_index_from_string, get_column_letter
from openpyxl.utils.datetime import to_excel

from .emitter import emit
from .engine import Engine
from .lexer import LexError
from .parser import ParseError, parse_formula
from .refs import MAX_COL, MAX_ROW
from .relocate import move_position, move_span
from .values import BLANK, ERROR_CODES, UNKNOWN, ErrorValue, is_error
from .workbook import MISSING, Cell, Workbook, formula_cell, literal_cell

# zip-part categories whose loss the fidelity guard detects
_PART_CATEGORIES = {
    "charts": r"^xl/charts/chart[^/]*\.xml$",
    "chart sheets": r"^xl/chartsheets/[^/]*\.xml$",
    "drawings": r"^xl/drawings/[^/]*\.xml$",
    "images/media": r"^xl/media/",
    "pivot tables": r"^xl/pivotTables/[^/]*\.xml$",
    "pivot caches": r"^xl/pivotCache/[^/]*\.xml$",
    "tables": r"^xl/tables/[^/]*\.xml$",
    "comments": r"^xl/comments[^/]*\.xml$",
    "threaded comments": r"^xl/threadedComments/",
    "VBA project": r"^xl/vbaProject\.bin$",
    "slicers": r"^xl/slicers/",
    "slicer caches": r"^xl/slicerCaches/",
    "timelines": r"^xl/timelines/",
    "external links": r"^xl/externalLinks/[^/]*\.xml$",
    "form controls": r"^xl/ctrlProps/",
    "ActiveX controls": r"^xl/activeX/",
    "embedded objects": r"^xl/embeddings/",
    "custom XML": r"^customXml/",
}


class WriteRefused(Exception):
    """apply could not write; the input file is untouched. Carries (code, message)."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code, self.message = code, message


@dataclass
class Loaded:
    path: Path
    model: Workbook
    engine: Engine
    xl: object                               # openpyxl Workbook (formulas)
    sha256: str
    load_warnings: list = field(default_factory=list)
    unparsed: int = 0                        # formulas Tabula cannot parse
    unparsed_names: int = 0


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def lock_file(path: Path) -> Path | None:
    """Excel's owner file (~$name) next to the workbook, if the workbook is open."""
    for candidate in (f"~${path.name}", f"~${path.name[2:]}"):
        p = path.with_name(candidate)
        if p.exists():
            return p
    return None


def _convert(v, data_type: str | None = None):
    """openpyxl value -> Tabula value."""
    if v is None:
        return BLANK
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, (datetime, date, time)):
        return float(to_excel(v))
    if isinstance(v, str) and data_type == "e":  # text that merely looks like "#N/A" stays text
        return ErrorValue(v)
    return v if isinstance(v, str) else str(v)


def load(path) -> Loaded:
    path = Path(path)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        keep_vba = path.suffix.lower() == ".xlsm"
        xl = openpyxl.load_workbook(path, keep_vba=keep_vba)
        cached_wb = openpyxl.load_workbook(path, data_only=True)
    model = Workbook()
    unparsed = 0
    for ws in xl.worksheets:
        sheet = model.add_sheet(ws.title)
        cws = cached_wb[ws.title]
        # Features holding cell addresses that Tabula does not relocate: structural
        # edits on this sheet are refused while they exist (E-STRUCT).
        sheet.meta = {
            "merged cell ranges": len(ws.merged_cells.ranges),
            "tables": len(ws.tables),
            "conditional formats": len(ws.conditional_formatting),
            "data validations": len(ws.data_validations.dataValidation),
            "hyperlinks": len(ws._hyperlinks) + sum(
                1 for c in getattr(ws, "_cells", {}).values() if getattr(c, "hyperlink", None)),
            "print areas/titles": int(bool(ws.print_area or ws.print_title_rows
                                           or ws.print_title_cols)),
            "autofilters": int(bool(ws.auto_filter.ref)),
        }
        # rule formulas may point at OTHER sheets; those sheets cannot be edited structurally
        sheet.rule_texts = [("data validations", text)
                            for dv in ws.data_validations.dataValidation
                            for text in (dv.formula1, dv.formula2) if text]
        sheet.rule_texts += [("conditional formats", text)
                             for cf in ws.conditional_formatting for rule in cf.rules
                             for text in (rule.formula or [])]
        # iterate stored cells only: a stray format at XFD1048576 must not cost 17bn cells
        stored = getattr(ws, "_cells", None)
        cells = list(stored.values()) if stored is not None else [c for r in ws.iter_rows() for c in r]
        for c in cells:
            v = c.value
            if v is None:
                continue
            if c.data_type == "f":
                cached_raw = cws.cell(row=c.row, column=c.column)
                cached = _convert(cached_raw.value, cached_raw.data_type)
                cached = MISSING if cached_raw.value is None else cached
                text = v if isinstance(v, str) else getattr(v, "text", None)
                cell = None
                if isinstance(v, str):
                    try:
                        cell = formula_cell(parse_formula(text), text, from_file=True)
                    except (LexError, ParseError):
                        cell = None
                if cell is None:  # array/data-table formula or syntax Tabula lacks
                    unparsed += 1
                    cell = Cell(raw=str(text or v), kind="unparsed",
                                value=UNKNOWN if cached is MISSING else cached)
                cell.cached = cached
            else:
                cell = literal_cell(_convert(v, c.data_type))
                if isinstance(v, str):
                    cell.raw = v
            sheet.cells[(c.column, c.row)] = cell
    unparsed_names = _load_names(xl, model)
    engine = Engine(model)
    engine.rebuild()
    engine.recompute(initial=True)
    msgs = sorted({str(w.message) for w in caught})
    return Loaded(path, model, engine, xl, sha256(path), msgs, unparsed, unparsed_names)


def _load_names(xl, model: Workbook) -> int:
    bad = 0

    def add(name, text, scope):
        nonlocal bad
        try:
            ast = parse_formula(text)
        except (LexError, ParseError):
            ast = None
            bad += 1
        model.define(name, ast, text, scope)

    for name, dn in xl.defined_names.items():
        if dn.attr_text:
            add(name, dn.attr_text, None)
    for ws in xl.worksheets:
        sid = model.sheet(ws.title).sid
        for name, dn in ws.defined_names.items():
            if dn.attr_text:
                add(name, dn.attr_text, sid)
    return bad


def part_counts(path: Path) -> Counter:
    counts: Counter = Counter()
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            for category, pattern in _PART_CATEGORIES.items():
                if re.match(pattern, name):
                    counts[category] += 1
    return counts


def losses(before: Counter, after: Counter) -> list[str]:
    return [f"{cat}: {before[cat]} -> {after.get(cat, 0)}"
            for cat in sorted(before) if after.get(cat, 0) < before[cat]]


def predict_loss(loaded: Loaded) -> list[str]:
    """What an unmodified load/save round trip would lose (used by `check`)."""
    before = part_counts(loaded.path)
    if not before and not loaded.load_warnings:
        return []
    fd, tmp = tempfile.mkstemp(suffix=loaded.path.suffix)
    os.close(fd)
    try:
        loaded.xl.save(tmp)
        lost = losses(before, part_counts(Path(tmp)))
    finally:
        os.unlink(tmp)
    return lost + [f"openpyxl: {m}" for m in loaded.load_warnings]


def _xl_value(cell: Cell):
    if cell.kind == "formula":
        return emit(cell.ast, plain_functions=FORMULAE)
    v = cell.value
    if v is BLANK:
        return None
    if is_error(v):
        return v.code
    if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
        return int(v)
    return v


def write_back(loaded: Loaded, structure: list[tuple], out_path, allow_lossy: bool = False) -> int:
    """Replay structural actions, sync changed cells, save atomically. Returns cells synced.

    structure: ordered tuples ('insert_rows'|'delete_rows'|'insert_cols'|'delete_cols',
    sheet, at, n), ('add_sheet', name) or ('rename_sheet', old, new).
    """
    out_path = Path(out_path)
    xl = loaded.xl
    try:
        for action in structure:
            kind = action[0]
            if kind == "add_sheet":
                xl.create_sheet(action[1])
            elif kind == "rename_sheet":
                ws = xl[action[1]]
                if action[1].casefold() == action[2].casefold():
                    ws.title = ".tabula-rename"  # openpyxl would de-duplicate 'SALES' vs 'Sales'
                ws.title = action[2]
            else:
                ws = xl[action[1]]
                getattr(ws, kind)(action[2], action[3])
                _shift_dimensions(ws, kind, action[2], action[3])
    except (KeyError, ValueError) as e:
        raise WriteRefused("E-WRITE", f"could not replay the structural edits: {e}")

    expected = {}
    for sheet in loaded.model.sheets:
        ws = xl[sheet.name]
        for (col, row), cell in sheet.cells.items():
            if not (cell.touched or cell.text_changed):
                continue
            target = ws.cell(row=row, column=col)
            value = _xl_value(cell)
            target.value = value
            if cell.kind != "formula" and isinstance(value, str) and (
                    value.startswith("=") or value in ERROR_CODES):
                target.data_type = "s"  # literal text, not a formula or an error value
            expected[(sheet.name, col, row)] = value
    for dn in loaded.model.names.values():
        if dn.changed and dn.ast is not None:
            text = emit(dn.ast, with_equals=False, plain_functions=FORMULAE)
            if dn.scope is None:
                xl.defined_names[dn.name].attr_text = text
            else:
                xl[loaded.model.by_sid(dn.scope).name].defined_names[dn.name].attr_text = text

    fd, tmp = tempfile.mkstemp(prefix=".tabula-", suffix=out_path.suffix,
                               dir=out_path.parent)
    os.close(fd)
    tmp = Path(tmp)
    try:
        xl.save(tmp)
        lost = losses(part_counts(loaded.path), part_counts(tmp))
        lost += [f"openpyxl: {m}" for m in loaded.load_warnings]
        if lost and not allow_lossy:
            raise WriteRefused("E-FIDELITY", "saving would lose workbook parts: " + "; ".join(lost))
        _verify(tmp, expected)
        os.replace(tmp, out_path)
    finally:
        if tmp.exists():
            tmp.unlink()
    return len(expected)


def _shift_dimensions(ws, kind: str, at: int, n: int) -> None:
    """openpyxl's insert/delete moves cells but not row heights, hidden rows or column
    widths; move those too so formatting stays with its data."""
    signed = n if kind.startswith("insert") else -n
    if kind.endswith("rows"):
        dims = dict(ws.row_dimensions)
        ws.row_dimensions.clear()
        for idx, dim in dims.items():
            q = move_position(idx, at, signed, MAX_ROW)
            if q is not None:
                dim.index = q
                ws.row_dimensions[q] = dim
        return
    dims = dict(ws.column_dimensions)
    ws.column_dimensions.clear()
    for letter, dim in dims.items():
        lo = dim.min or column_index_from_string(letter)
        hi = dim.max or lo
        span = move_span(lo, hi, at, signed, MAX_COL)
        if span is None:
            continue
        dim.min, dim.max = span
        dim.index = get_column_letter(span[0])
        ws.column_dimensions[dim.index] = dim


def _verify(path: Path, expected: dict) -> None:
    """Reload the written file and check every synced cell holds what we meant to write."""
    check = openpyxl.load_workbook(path)
    for (sheet_name, col, row), value in expected.items():
        got = check[sheet_name].cell(row=row, column=col).value
        same = got == value or (isinstance(got, (int, float)) and isinstance(value, (int, float))
                                and abs(got - value) <= 1e-12 * max(1.0, abs(value)))
        if not same:
            raise WriteRefused("E-VERIFY", f"post-write verification failed at {sheet_name}!"
                               f"R{row}C{col}: wrote {value!r}, read back {got!r}")
