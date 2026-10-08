"""T1 with plain openpyxl (scripted reference for the default arm; correct)."""
import openpyxl

wb = openpyxl.load_workbook("book.xlsx")
ws = wb["Sales"]
ws["G1"] = "Tax"
for r in range(2, 22):
    ws[f"G{r}"] = f"=E{r}*Inputs!$B$2"
ws["G22"] = "=SUM(G2:G21)"
wb["Summary"]["A8"] = "Total tax"
wb["Summary"]["B8"] = "=Sales!G22"
wb.save("book.xlsx")
print("added Tax column G and Summary!B8")
