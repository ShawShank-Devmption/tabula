"""T3 with plain openpyxl: the straightforward approach (scripted reference, naive)."""
import openpyxl

wb = openpyxl.load_workbook("book.xlsx")
ws = wb["Sales"]
for r in sorted((r for r in range(2, ws.max_row + 1) if ws[f"B{r}"].value == "TEST"), reverse=True):
    ws.delete_rows(r)
wb.save("book.xlsx")
print("deleted TEST rows")
