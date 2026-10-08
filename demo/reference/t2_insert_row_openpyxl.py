"""T2 with plain openpyxl: the straightforward approach (scripted reference, naive).

ws.insert_rows() moves cells but does not rewrite formulas or defined names.
"""
import openpyxl

wb = openpyxl.load_workbook("book.xlsx")
ws = wb["Sales"]
ws.insert_rows(11)
for col, value in zip("ABCD", ["West", "Dana Lee", 40, 125]):
    ws[f"{col}11"] = value
ws["E11"] = "=C11*D11"
ws["F11"] = "=E11*CommissionRate"
wb.save("book.xlsx")
print("inserted row 11")
