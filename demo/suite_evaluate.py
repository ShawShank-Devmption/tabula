"""Independent held-out checker: plain Python arithmetic, pycel and preservation.

This file does not import Tabula. Missing/partial run manifests are errors.
"""
from __future__ import annotations
import argparse
import json
import logging
import math
import re
import tempfile
from collections import Counter
from pathlib import Path

import openpyxl
from suite import TASKS_BY_ID, SPECS, SCORING_VERSION, build, solve, expectations, positions, names, sha256

logging.disable(logging.CRITICAL)
OUTCOMES=('successful_completion','safe_refusal','unverified','failure','infrastructure')


def final_status(message):
    matches=re.findall(r'^EVAL_STATUS: (SUCCESS|REFUSAL|UNVERIFIED|FAILURE)\s*$',message or '',re.M)
    return matches[0] if len(matches)==1 else None


def classify(meta,checks,refusal_task=False):
    claim=final_status(meta.get('final_message',''))
    valid=bool(checks) and all(c['ok'] for c in checks)
    infrastructure=meta.get('timed_out') or meta.get('runner_error') or meta.get('returncode') not in (0,) or meta.get('is_error')
    if infrastructure: outcome='infrastructure'
    elif claim=='SUCCESS': outcome='successful_completion' if valid and not refusal_task else 'failure'
    elif claim=='REFUSAL': outcome='safe_refusal' if valid else 'failure'
    elif claim=='FAILURE': outcome='failure'
    else:outcome='unverified'
    return dict(outcome=outcome,claimed_status=claim,workbook_valid=valid,
                silent_corruption=bool(claim=='SUCCESS' and not valid and not infrastructure))


def cell_style(cell):
    # Compare semantic objects, never style ids (ids differ across serializers).
    return (str(cell.font),str(cell.fill),str(cell.border),str(cell.alignment),cell.number_format,str(cell.protection))


def cells(ws):
    return {c.coordinate:c for row in ws.iter_rows() for c in row if c.value is not None or c.has_style}


def schema(wb):
    return dict(sheets=wb.sheetnames,names=sorted((n.name,n.attr_text,n.localSheetId,n.hidden) for n in names(wb)),epoch=str(wb.epoch),
        worksheets={ws.title:dict(state=ws.sheet_state,freeze=str(ws.freeze_panes),merged=sorted(str(r) for r in ws.merged_cells.ranges),
            columns={k:(v.width,v.hidden,v.outlineLevel) for k,v in ws.column_dimensions.items()},
            rows={k:(v.height,v.hidden,v.outlineLevel) for k,v in ws.row_dimensions.items()},
            autofilter=ws.auto_filter.ref,print_area=str(ws.print_area),print_titles=str(ws.print_title_rows),
            tables=sorted((t.name,t.ref) for t in ws.tables.values()),charts=[str(c.to_tree()) for c in ws._charts],
            validations=str(ws.data_validations),conditional=str(list(ws.conditional_formatting))) for ws in wb.worksheets})


def check_task(task_id,path,baseline=None):
    task=TASKS_BY_ID[task_id]; s=SPECS[task.workbook]; p=positions(task)
    path=Path(path); checks=[]
    def add(name,ok,detail=''):
        checks.append(dict(check=name,ok=bool(ok),detail='' if ok else str(detail)[:1800]))
    try:wb=openpyxl.load_workbook(path)
    except Exception as exc:
        add('workbook opens',False,exc);return checks
    with tempfile.TemporaryDirectory() as tmp:
        original=Path(tmp)/'expected.xlsx'
        build(task,original);solve(task,original)
        expected=openpyxl.load_workbook(original)
        # Chart XML object repr contains addresses: compare source XML by serialized bytes instead.
        from openpyxl.xml.functions import tostring
        sa,sb=schema(wb),schema(expected)
        for sw,book in ((sa,wb),(sb,expected)):
            for ws in book.worksheets:sw['worksheets'][ws.title]['charts']=[tostring(c.to_tree()).decode() for c in ws._charts]
        add('schema, names and workbook features preserved',sa==sb, f'actual schema: {sa}' if sa!=sb else '')
        bad=[]
        expected_values=expectations(task)
        for ws in expected.worksheets:
            if ws.title not in wb.sheetnames:
                bad.append(f'missing {ws.title}');continue
            a,b=cells(wb[ws.title]),cells(ws)
            for addr in set(a)|set(b):
                actual=a.get(addr); want=b.get(addr)
                av=actual.value if actual else None; ev=want.value if want else None
                formula=isinstance(ev,str) and ev.startswith('=')
                # All edited/relocated formula outputs are scored semantically; preserved
                # formula expressions elsewhere must remain byte-identical.
                flexible=task.operation in ('insert','delete','rename') or (task.operation=='fix' and ws.title==s.summary and addr in (f'B{s.summary_row}',f'B{s.summary_row+1}')) or (task.operation=='add' and ((ws.title==s.sheet and addr.startswith('F')) or (ws.title==s.summary and addr==f'B{s.summary_row+3}')))
                if formula:
                    if not isinstance(av,str) or not av.startswith('=') or (not flexible and av!=ev):bad.append(f'{ws.title}!{addr}: formula changed or replaced')
                elif av!=ev:bad.append(f'{ws.title}!{addr}: {av!r} != {ev!r}')
                # Formatting on cells introduced by the task is intentionally unrestricted.
                new=(task.operation=='insert' and ws.title==s.sheet and actual is not None and actual.row==s.header+3) or (task.operation=='add' and ((ws.title==s.sheet and addr.startswith('F')) or (ws.title==s.summary and addr in (f'A{s.summary_row+3}',f'B{s.summary_row+3}'))))
                if not new and actual and want and cell_style(actual)!=cell_style(want):bad.append(f'{ws.title}!{addr}: style changed')
                if actual and want and (str(actual.comment)!=str(want.comment) or str(actual.hyperlink)!=str(want.hyperlink)):bad.append(f'{ws.title}!{addr}: annotation changed')
        add('cell values, formulas and styles preservation',not bad,'; '.join(bad[:12]))
        if task.operation.startswith('refuse') and baseline is not None:
            add('safe refusal leaves workbook byte-for-byte unchanged',sha256(path)==sha256(baseline),'output hash differs from frozen input')
        try:
            from pycel import ExcelCompiler
            compiler=ExcelCompiler(filename=str(path),cycles=False)
            def evaluate(oracle,xl):
                failed=[]
                for (sheet,addr),want in oracle.items():
                    try:
                        got=xl.evaluate(f"'{sheet}'!{addr}")
                        if isinstance(got,bool) or not isinstance(got,(int,float)) or not math.isclose(got,want,rel_tol=1e-8,abs_tol=1e-7):failed.append(f'{sheet}!{addr}: {got!r} != {want}')
                    except Exception as exc:failed.append(f'{sheet}!{addr}: {type(exc).__name__}')
                return failed
            errors=evaluate(expected_values,compiler)
            add('all formula values match plain Python oracle',not errors,'; '.join(errors[:8]))
            errors=[]
            for ws in wb.worksheets:
                for c in cells(ws).values():
                    if c.data_type=='f':
                        try:
                            v=compiler.evaluate(f"'{ws.title}'!{c.coordinate}")
                            if isinstance(v,str) and v.startswith('#'):errors.append(f'{ws.title}!{c.coordinate}: {v}')
                        except Exception as exc:errors.append(f'{ws.title}!{c.coordinate}: {type(exc).__name__}')
            add('no formula errors or cycles anywhere',not errors,'; '.join(errors[:8]))
            probes=[(('Settings',s.rate_cell),0.13)]
            for i,row in enumerate(p['rows']):
                r=p['first']+i
                probes.extend([((p['sheet'],f'{p["cols"]["qty"]}{r}'),row[2]+7),((p['sheet'],f'{p["cols"]["price"]}{r}'),row[3]+13)])
            errors=[]
            for (sheet,addr),new in probes:
                xl=ExcelCompiler(filename=str(path),cycles=False)
                evaluate(expected_values,xl)  # register input dependencies before set_value
                try:
                    xl.set_value(f"'{sheet}'!{addr}",new)
                    failures=evaluate(expectations(task,{(sheet,addr):new}),xl)
                    errors.extend(f'{sheet}!{addr} probe: {f}' for f in failures)
                except Exception as exc:errors.append(f'{sheet}!{addr} probe: {type(exc).__name__}')
            add('liveness under every quantity, price and rate probe',not errors,'; '.join(errors[:8]))
        except ImportError as exc:
            raise RuntimeError('Independent checker requires demo/.venv pycel; cannot score without it') from exc
        except Exception as exc:
            add('independent evaluator completed',False,f'{type(exc).__name__}: {exc}')
    return checks


def verify_manifest(folder):
    folder=Path(folder)
    manifest=json.loads((folder/'manifest.json').read_text())
    if manifest['scoring_version']!=SCORING_VERSION:raise ValueError('scoring version mismatch')
    if manifest['checker_sha256']!=sha256(__file__):raise ValueError('checker changed since run freeze')
    for task in manifest['tasks']:
        if sha256(folder/'fixtures'/task['fixture'])!=task['sha256']:raise ValueError(f'fixture hash mismatch: {task["id"]}')
    dirs={j['directory'] for j in manifest['jobs']}
    found={p.parent.name for p in folder.glob('*/meta.json')}
    if dirs!=found:raise ValueError(f'partial or unexpected runs; missing={sorted(dirs-found)}, extra={sorted(found-dirs)}')
    return manifest


def evaluate_runs(folder):
    folder=Path(folder);manifest=verify_manifest(folder);results=[]
    fixture={t['id']:t for t in manifest['tasks']}
    for job in manifest['jobs']:
        rundir=folder/job['directory'];meta=json.loads((rundir/'meta.json').read_text())
        if (meta.get('task'),meta.get('arm'),meta.get('trial'))!=(job['task'],job['arm'],job['trial']):raise ValueError('run identity mismatch')
        checks=check_task(job['task'],rundir/'book.xlsx',folder/'fixtures'/fixture[job['task']]['fixture'])
        meta.update(checks=checks,passed=sum(c['ok'] for c in checks),total=len(checks))
        meta.update(classify(meta,checks,TASKS_BY_ID[job['task']].operation.startswith('refuse')))
        results.append(meta)
    (folder/'results.json').write_text(json.dumps(results,indent=2)+'\n')
    write_report(folder,manifest,results)
    return results


def write_report(folder,manifest,results):
    lines=['# Controlled held-out workbook evaluation','',f"Suite `{manifest['suite_version']}`; scoring `{SCORING_VERSION}`.",
      f"Driver: {manifest['driver']}; model requested: `{manifest['model']}`; runner version: `{manifest['runner_version']}`.",
      'Independent checks use pycel plus plain Python; Tabula is never the grading engine.',
      'Reference results are scripted calibration, not evidence of agent performance.','',
      '| Arm | Successful completion | Safe refusal | Unverified | Failure | Infrastructure | Claimed success with corrupt workbook | Time total (s) | Estimated USD total | Tool calls | Tabula used |',
      '|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|']
    for arm in manifest['arms']:
        rs=[r for r in results if r['arm']==arm];counts=Counter(r['outcome'] for r in rs)
        sums=[sum(r.get(k) or 0 for r in rs) for k in ('duration_s','cost_usd','tool_calls')]
        cost=f'{sums[1]:.4f}' if any(r.get('cost_usd') is not None for r in rs) else 'unavailable'
        lines.append('| '+arm+' | '+' | '.join(str(counts[o]) for o in OUTCOMES)+f" | {sum(r['silent_corruption'] for r in rs)} | {sums[0]:.2f} | {cost} | {sums[2]} | {sum(bool(r.get('tabula_used')) for r in rs)} |")
    lines+=['','Successful-completion denominator excludes the two refusal scenarios; refusal counts are reported separately. Silent corruption requires an explicit SUCCESS status and a failed workbook check. Normal process exit alone is unverified.',
      'Costs are Claude CLI estimates at API list prices, not subscription charges. Full model usage and resolved model ids are retained in results.json.','',
      '| Task | Trial | Default outcome | Tabula outcome | Default / Tabula time (s) | Default / Tabula cost (USD) | Default / Tabula tool calls |',
      '|---|---:|---|---|---|---|---|']
    keys=sorted({(r['task'],r['trial']) for r in results})
    for task,trial in keys:
        pair={r['arm']:r for r in results if r['task']==task and r['trial']==trial}
        d,t=pair.get('default',{}),pair.get('tabula',{})
        lines.append(f"| {task} | {trial} | {d.get('outcome','missing')} | {t.get('outcome','missing')} | {d.get('duration_s','—')} / {t.get('duration_s','—')} | {d.get('cost_usd','—')} / {t.get('cost_usd','—')} | {d.get('tool_calls','—')} / {t.get('tool_calls','—')} |")
    lines+=['','## Failed checks','']
    for r in results:
        failed=[c for c in r['checks'] if not c['ok']]
        if failed:
            lines.append(f"- `{r['task']} {r['arm']} trial {r['trial']}` ({r['outcome']}): "+'; '.join(c['check']+': '+c['detail'].replace('\n',' ') for c in failed))
    if not any(not c['ok'] for r in results for c in r['checks']):lines.append('No failed workbook checks.')
    lines+=['','## Limits','',
      '- These fifteen synthetic tasks are held out from the five original development demos. They are not a representative sample of every Excel feature.',
      '- One model and three trials do not establish generality or statistical significance; task pairs and raw transcripts support inspection.',
      '- Python/openpyxl permissions, common workflow advice, prompts, budgets and timeouts are matched. The Tabula arm additionally gets TEL.md and the Tabula wrapper; actual usage is reported.',
      '- No native Excel recalculation is available. pycel covers the simple arithmetic, SUM and SUMIF used here; these checks do not prove general Excel equivalence.',
      '- CLI allowlists and safe-mode isolate configuration, not filesystem/network access. Agents are instructed to stay in their run folder; reference fixtures and grader are not disclosed. Source/workbook hashes and transcripts make deviations inspectable.','']
    (folder/'REPORT.md').write_text('\n'.join(lines))

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('folder',type=Path)
    args=ap.parse_args()
    results=evaluate_runs(args.folder)
    print(json.dumps(dict(Counter(r['outcome'] for r in results)),sort_keys=True))
    print(f'Wrote {args.folder}/REPORT.md and results.json')
