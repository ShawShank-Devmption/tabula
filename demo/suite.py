"""Frozen held-out scenarios and independent Python oracle (no Tabula imports).

Five layouts, fifteen tasks. Each task starts from a fresh task-specific fixture.
Fixture ZIP timestamps and workbook document properties are normalized for stable hashes.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from copy import copy
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
SUITE_VERSION = 'holdout-2026-10-08-v1'
SCORING_VERSION = 'preservation-python-pycel-liveness-v1'

@dataclass(frozen=True)
class Spec:
    sheet: str
    summary: str
    header: int
    columns: tuple[str, ...]
    labels: tuple[str, ...]
    rate_cell: str
    summary_row: int
    rows: tuple[tuple, ...]  # id, group, quantity, unit cost

SPECS = {
 'inventory': Spec('Stock Ledger', 'Warehouse Summary', 1, ('id','group','qty','price','amount'), ('SKU','Warehouse','Count','Unit cost','Value'), 'B2', 3,
  (('IV-18','East',17,11),('IV-02','West',8,23),('IV-91','East',6,47),('IV-44','TEST',1,2),('IV-76','West',13,19),('IV-30','East',4,61))),
 'payroll': Spec('Hours', 'Payroll Review', 3, ('group','id','price','qty','amount'), ('Team','Staff code','Hourly rate','Hours','Pay'), 'C4', 5,
  (('PY-23','East',14,32),('PY-09','West',20,29),('PY-88','East',9,51),('PY-17','West',12,37),('PY-64','East',18,26))),
 'orders': Spec('Order Lines', 'Order Overview', 2, ('id','qty','group','amount','price'), ('Order id','Quantity','Channel','Revenue','Price'), 'D3', 2,
  (('OR-53','West',3,149),('OR-12','East',11,27),('OR-74','TEST',1,3),('OR-29','East',7,43),('OR-06','West',5,81),('OR-38','TEST',2,4),('OR-95','East',6,67))),
 'grants': Spec('Grant Budget', 'Board Dashboard', 4, ('price','group','id','qty','amount'), ('Cost','Programme','Grant id','Awards','Budget'), 'B5', 4,
  (('GR-82','East',5,211),('GR-13','West',9,137),('GR-57','East',7,173),('GR-40','West',4,251),('GR-66','East',8,109),('GR-04','West',3,317),('GR-31','East',6,193),('GR-99','West',2,401))),
 'fleet': Spec('Vehicle Costs', 'Fleet Totals', 2, ('qty','price','id','group','amount'), ('Trips','Cost per trip','Vehicle id','Depot','Spend'), 'C2', 6,
  (('FL-71','East',12,31),('FL-08','West',7,53),('FL-42','East',19,17),('FL-90','West',4,89))),
}

@dataclass(frozen=True)
class Task:
    id: str
    workbook: str
    operation: str
    prompt: str


def _task(id, workbook, operation, request):
    s = SPECS[workbook]
    return Task(id, workbook, operation, request + ' Preserve all unrelated records, formulas, formatting, sheets, defined names and workbook features. Keep calculations as live formulas; all existing summaries must remain correct.')

TASKS = [
 _task('e01_inventory_fee','inventory','add', 'On Stock Ledger, add column F headed "Service fee", with Value multiplied by the service rate on Settings for each item and a live total. Add "Total service fee" to Warehouse Summary below its existing figures, linking its live total.'),
 _task('e02_inventory_insert','inventory','insert', 'Insert a missing item immediately before SKU IV-91 on Stock Ledger: SKU IV-NEW, warehouse West, count 9, unit cost 37, with a live Value formula. Include it in every total and summary.'),
 _task('e03_inventory_delete','inventory','delete', 'Delete every entire record row marked TEST in the Warehouse column of Stock Ledger. Keep totals, summaries and names correct after rows move.'),
 _task('e04_payroll_rename','payroll','rename', 'Rename Hours to "Approved Hours" and keep every cross-sheet reference and defined name working.'),
 _task('e05_payroll_fix','payroll','fix', 'Two formulas on Payroll Review are wrong: the total and East figure. Locate and repair them to cover exactly all records and the correct team. Do not modify other sheets.'),
 _task('e06_payroll_identify','payroll','identify', 'Use Directory to identify the staff code for "Mira Santos"; then change only that person\'s Hours quantity to 19. The main sheet contains codes, not display names.'),
 _task('e07_orders_fee','orders','add', 'On Order Lines, add column F headed "Service fee", with Revenue multiplied by the service rate on Settings for each item and a live total. Add "Total service fee" below the existing Order Overview figures, linking its live total.'),
 _task('e08_orders_delete','orders','delete', 'Delete every entire record row marked TEST in the Channel column of Order Lines. Keep totals, summaries and names correct after rows move.'),
 _task('e09_orders_identify','orders','identify', 'Use Directory to identify the order id for "Mira Santos"; change only that order\'s Quantity to 19. Preserve the other order records.'),
 _task('e10_grants_insert','grants','insert', 'Insert a grant immediately before GR-57 on Grant Budget: grant id GR-NEW, programme West, awards 9, cost 37, with a live Budget formula. Include it in totals and summaries.'),
 _task('e11_grants_rename','grants','rename', 'Rename Grant Budget to "Approved Grants" and keep every cross-sheet reference and defined name working.'),
 _task('e12_grants_fix','grants','fix', 'Two formulas on Board Dashboard are wrong: total budget and East figure. Locate and repair them to cover exactly all records and the correct programme. Do not modify other sheets.'),
 _task('e13_fleet_fix','fleet','fix', 'Two formulas on Fleet Totals are wrong: total spend and East figure. Locate and repair them to cover exactly all records and the correct depot. Do not modify other sheets.'),
 _task('e14_fleet_chart_refusal','fleet','refuse_chart', 'Vehicle Costs has an embedded chart. A proposed insertion would move its source range. This task requires a certified automatic relocation of all chart references; manual edits or stripping the chart are forbidden. If your tooling does not provide that guarantee, leave the workbook byte-for-byte unchanged and report a safe refusal explaining the unsupported chart relocation.'),
 _task('e15_orders_table_refusal','orders','refuse_table', 'Order Lines has an Excel table. A proposed row deletion would alter its range. This task requires a certified automatic relocation of all table metadata; manual edits or stripping the table are forbidden. If your tooling does not provide that guarantee, leave the workbook byte-for-byte unchanged and report a safe refusal explaining the unsupported table relocation.'),
]
TASKS_BY_ID = {t.id: t for t in TASKS}


def names(wb):
    container = wb.defined_names
    return list(container.values()) if hasattr(container,'values') else list(container.definedName)


def put_name(wb, name, ref):
    defined = DefinedName(name, attr_text=ref)
    if hasattr(wb.defined_names, 'add'):
        wb.defined_names.add(defined)
    else:
        wb.defined_names.append(defined)


def positions(task, solved=True):
    s = SPECS[task.workbook]
    rows = list(s.rows)
    if solved and task.operation == 'insert':
        rows.insert(2, (s.rows[0][0][:2] + '-NEW','West',9,37))
    if solved and task.operation == 'delete':
        rows = [r for r in rows if r[1] != 'TEST']
    if solved and task.operation == 'identify':
        rows[1] = (rows[1][0], rows[1][1],19,rows[1][3])
    first = s.header + 1
    sheet = ('Approved Hours' if task.workbook == 'payroll' else 'Approved Grants') if solved and task.operation == 'rename' else s.sheet
    cs = {k: col(i+1) for i,k in enumerate(s.columns)}
    return dict(rows=rows, first=first,last=first+len(rows)-1,total=first+len(rows),sheet=sheet,cols=cs,
                amount_cells=[f"{cs['amount']}{first+i}" for i in range(len(rows))])


def write_formulas(wb, task, solved):
    s = SPECS[task.workbook]
    p = positions(task,solved)
    ws = wb[p['sheet']]
    c = p['cols']
    for i,_ in enumerate(p['rows']):
        r=p['first']+i
        ws[f"{c['amount']}{r}"] = f"={c['qty']}{r}*{c['price']}{r}"
    t,first,last = p['total'],p['first'],p['last']
    ws[f"{c['amount']}{t}"] = f"=SUM({c['amount']}{first}:{c['amount']}{last})"
    ws[f"{c['qty']}{t}"] = f"=SUM({c['qty']}{first}:{c['qty']}{last})"
    sr=s.summary_row
    summary=wb[s.summary]
    total_last=last-1 if task.operation=='fix' and not solved else last
    summary[f'B{sr}']=f"=SUM('{p['sheet']}'!{c['amount']}{first}:{c['amount']}{total_last})"
    group=c['id'] if task.operation=='fix' and not solved else c['group']
    summary[f'B{sr+1}']=f'=SUMIF(\'{p["sheet"]}\'!{group}{first}:{group}{last},"East",\'{p["sheet"]}\'!{c["amount"]}{first}:{c["amount"]}{last})'
    summary[f'B{sr+2}']=f"=B{sr}*Settings!${s.rate_cell[0]}${s.rate_cell[1:]}"
    for dn in names(wb):
        if dn.name=='Amounts':
            dn.attr_text=f"'{p['sheet']}'!${c['amount']}${first}:${c['amount']}${last}"
    if solved and task.operation=='add':
        ws.cell(s.header,6,'Service fee')
        for r in range(first,last+1):
            ws.cell(r,6,f"={c['amount']}{r}*Settings!${s.rate_cell[0]}${s.rate_cell[1:]}")
        ws.cell(t,6,f'=SUM(F{first}:F{last})')
        summary.cell(sr+3,1,'Total service fee')
        summary.cell(sr+3,2,f"='{p['sheet']}'!F{t}")


def stable_save(wb,path):
    path=Path(path)
    fixed=datetime(2026,10,8)
    wb.properties.created=wb.properties.modified=fixed
    wb.save(path)
    with zipfile.ZipFile(path) as src:
        content={n:src.read(n) for n in src.namelist()}
    # openpyxl 3.1 stamps modified on save; normalize its XML as well.
    content['docProps/core.xml']=re.sub(rb'<dcterms:modified[^>]*>.*?</dcterms:modified>',b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:type="dcterms:W3CDTF">2026-10-08T00:00:00Z</dcterms:modified>',content['docProps/core.xml'])
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(content):
            info=zipfile.ZipInfo(name,date_time=(2026,10,8,0,0,0))
            info.compress_type=zipfile.ZIP_DEFLATED
            dst.writestr(info,content[name])


def build(task,path):
    s=SPECS[task.workbook]
    p=positions(task,False)
    wb=Workbook()
    settings=wb.active
    settings.title='Settings'
    settings['A1']='Evaluation settings'
    settings[s.rate_cell]=0.07
    settings['F7']='Do not edit this note'
    ws=wb.create_sheet(s.sheet)
    if s.header>1:
        ws['A1']=f'{task.workbook.title()} FY2026'
        ws['A1'].font=Font(bold=True,color='24557A')
    for i,label in enumerate(s.labels,1):
        cell=ws.cell(s.header,i,label)
        cell.font=Font(bold=True)
        cell.fill=PatternFill('solid',fgColor='D8E9EE')
        ws.column_dimensions[col(i)].width=18+i
    ws.freeze_panes=f'A{s.header+1}'
    for i,row in enumerate(s.rows):
        record=dict(zip(('id','group','qty','price'),row))
        for j,key in enumerate(s.columns,1):
            if key!='amount': ws.cell(p['first']+i,j,record[key])
        ws[f"{p['cols']['price']}{p['first']+i}"].number_format='0.00'
        ws[f"{p['cols']['amount']}{p['first']+i}"].number_format='#,##0.00'
    ws.cell(p['total'],1,'TOTAL')
    ws.cell(p['total'],1).font=Font(bold=True)
    summary=wb.create_sheet(s.summary)
    summary['A1']=f'{task.workbook.title()} report'
    for i,label in enumerate(('Total amount','East amount','Service fee estimate')):
        summary.cell(s.summary_row+i,1,label)
    summary.column_dimensions['A'].width=25
    directory=wb.create_sheet('Directory')
    directory.append(['Code','Display name','Annotation'])
    for i,row in enumerate(s.rows):
        directory.append([row[0],'Mira Santos' if i==1 else f'Contact {i+1}',f'Retain annotation {i+1}'])
    put_name(wb,'Amounts',f"'{s.sheet}'!${p['cols']['amount']}${p['first']}:${p['cols']['amount']}${p['last']}")
    put_name(wb,'ServiceRate',f"'Settings'!${s.rate_cell[0]}${s.rate_cell[1:]}")
    write_formulas(wb,task,False)
    if task.operation=='refuse_chart':
        chart=BarChart()
        chart.title='Vehicle spend'
        chart.add_data(Reference(ws,min_col=5,min_row=p['first'],max_row=p['last']))
        ws.add_chart(chart,'H2')
    if task.operation=='refuse_table':
        table=Table(displayName='ProtectedOrders',ref=f"A{s.header}:E{p['last']}")
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium9',showRowStripes=True)
        ws.add_table(table)
    stable_save(wb,path)


def solve(task,path):
    """Correct openpyxl calibration solution; fixture-specific relocation is explicit."""
    if task.operation.startswith('refuse'): return
    s=SPECS[task.workbook]
    wb=load_workbook(path)
    ws=wb[s.sheet]
    if task.operation=='insert':
        row=s.header+3
        ws.insert_rows(row)
        rec=dict(zip(('id','group','qty','price'),positions(task)['rows'][2]))
        for j,key in enumerate(s.columns,1):
            if key!='amount': ws.cell(row,j,rec[key])
    elif task.operation=='delete':
        for i in range(len(s.rows)-1,-1,-1):
            if s.rows[i][1]=='TEST':ws.delete_rows(s.header+1+i)
    elif task.operation=='rename':
        ws.title=positions(task)['sheet']
    elif task.operation=='identify':
        ws[f"{positions(task)['cols']['qty']}{s.header+2}"]=19
    write_formulas(wb,task,True)
    stable_save(wb,path)


def tel_solution(task):
    s=SPECS[task.workbook]
    p=positions(task)
    c=p['cols']; lines=[]
    if task.operation.startswith('refuse'):
        return ''
    if task.operation=='rename':
        return f'rename sheet "{s.sheet}" to "{p["sheet"]}"\n'
    lines.append(f'in "{s.sheet}" {{')
    if task.operation=='insert':
        row=s.header+3
        lines.append(f'insert rows {row}')
        record=dict(zip(('id','group','qty','price'),p['rows'][2]))
        for j,key in enumerate(s.columns,1):
            if key!='amount':lines.append(f'set {col(j)}{row} {json.dumps(record[key])}')
        lines.append(f'set {c["amount"]}{row} ={c["qty"]}{row}*{c["price"]}{row}')
    elif task.operation=='delete':
        for i in range(len(s.rows)-1,-1,-1):
            if s.rows[i][1]=='TEST':lines.append(f'delete rows {s.header+1+i}')
    elif task.operation=='identify':lines.append(f'set {c["qty"]}{s.header+2} 19')
    elif task.operation=='add':
        lines.extend([f'set F{s.header} "Service fee"',f'set F{p["first"]}:F{p["last"]} ={c["amount"]}{p["first"]}*Settings!${s.rate_cell[0]}${s.rate_cell[1:]}',f'set F{p["total"]} =SUM(F{p["first"]}:F{p["last"]})'])
    lines.append('}')
    if task.operation=='fix':
        lines.extend([f'in "{s.summary}" {{', f'set B{s.summary_row} =SUM(\'{s.sheet}\'!{c["amount"]}{p["first"]}:{c["amount"]}{p["last"]})',f'set B{s.summary_row+1} =SUMIF(\'{s.sheet}\'!{c["group"]}{p["first"]}:{c["group"]}{p["last"]},"East",\'{s.sheet}\'!{c["amount"]}{p["first"]}:{c["amount"]}{p["last"]})','}'])
    if task.operation=='add':lines.extend([f'in "{s.summary}" {{',f'set A{s.summary_row+3} "Total service fee"',f'set B{s.summary_row+3} =\'{s.sheet}\'!F{p["total"]}','}'])
    return '\n'.join(lines)+'\n'


def expectations(task,changed=None):
    """Plain arithmetic oracle. changed maps (sheet,address) to probe inputs."""
    s=SPECS[task.workbook]; p=positions(task); c=p['cols']; changed=changed or {}
    rate=changed.get(('Settings',s.rate_cell),0.07)
    result={}; revenues=[]; quantities=[]
    for i,(_,group,qty,price) in enumerate(p['rows']):
        r=p['first']+i
        qty=changed.get((p['sheet'],f'{c["qty"]}{r}'),qty)
        price=changed.get((p['sheet'],f'{c["price"]}{r}'),price)
        value=qty*price; revenues.append(value); quantities.append(qty)
        result[(p['sheet'],f'{c["amount"]}{r}')]=value
        if task.operation=='add':result[(p['sheet'],f'F{r}')]=value*rate
    result[(p['sheet'],f'{c["amount"]}{p["total"]}')]=sum(revenues)
    result[(p['sheet'],f'{c["qty"]}{p["total"]}')]=sum(quantities)
    sr=s.summary_row
    result[(s.summary,f'B{sr}')]=sum(revenues)
    result[(s.summary,f'B{sr+1}')]=sum(v for v,row in zip(revenues,p['rows']) if row[1]=='East')
    result[(s.summary,f'B{sr+2}')]=sum(revenues)*rate
    if task.operation=='add':
        result[(p['sheet'],f'F{p["total"]}')]=sum(revenues)*rate
        result[(s.summary,f'B{sr+3}')]=sum(revenues)*rate
    return result


def sha256(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def freeze(folder):
    folder=Path(folder)
    folder.mkdir(parents=True,exist_ok=False)
    data={'suite_version':SUITE_VERSION,'scoring_version':SCORING_VERSION,'tasks':[]}
    for task in TASKS:
        path=folder/f'{task.id}.xlsx'
        build(task,path)
        data['tasks'].append(dict(asdict(task),sha256=sha256(path),fixture=path.name))
    data['definition_sha256']=sha256(__file__)
    (folder/'scenarios.json').write_text(json.dumps(data,indent=2)+'\n')
    return data

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--freeze',required=True,type=Path)
    args=ap.parse_args()
    freeze(args.freeze)
    print(f'Frozen {len(TASKS)} scenarios in {args.freeze}')
