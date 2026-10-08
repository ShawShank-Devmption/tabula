"""Build the demo workbooks (design.md section 10).

    workbooks/sales_q3.xlsx        Inputs / Sales / Summary with defined names and
                                   cross-sheet formulas -- the starting point for T1, T2, T3, T5
    workbooks/sales_q3_buggy.xlsx  the same with two seeded Summary bugs -- for T4

DATA below is the single source of truth: the task checkers compute every
expected number from it in plain Python, independently of any formula engine.
Run:  python3 demo/build_workbook.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
OUT = HERE / "workbooks"

TAX_RATE = 0.18
COMMISSION_RATE = 0.06
TARGET = 60000

# region, rep, units, unit price.  Rows 8 and 15 are test entries that T3 removes.
# The last row is North on purpose: a SUMIF whose range stops one row short loses it.
DATA = [
    ("North", "Asha Rao", 42, 120),
    ("South", "Ben Ortiz", 18, 250),
    ("East", "Chen Wei", 35, 90),
    ("West", "Divya Nair", 27, 180),
    ("North", "Eli Cohen", 50, 75),
    ("South", "Farah Khan", 12, 300),
    ("Test", "TEST", 1, 1),
    ("East", "Gita Iyer", 64, 55),
    ("West", "Hugo Lind", 22, 210),
    ("North", "Ivan Petrov", 31, 140),
    ("South", "Jia Liu", 45, 95),
    ("East", "Kofi Mensah", 19, 260),
    ("West", "Lena Fischer", 38, 115),
    ("Test", "TEST", 1, 1),
    ("North", "Mateo Silva", 26, 175),
    ("South", "Nadia Haddad", 53, 68),
    ("East", "Omar Aziz", 29, 150),
    ("West", "Priya Menon", 41, 102),
    ("South", "Quinn Murphy", 16, 280),
    ("North", "Ravi Shah", 47, 88),
]
FIRST, LAST = 2, 1 + len(DATA)   # data rows 2..21
TOTAL_ROW = LAST + 1             # 22

HEAD = Font(bold=True)
FILL = PatternFill("solid", fgColor="E8EEF7")


def build(path: Path, buggy: bool = False) -> None:
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    for r, (label, value) in enumerate([("Assumption", "Value"), ("Tax rate", TAX_RATE),
                                        ("Commission rate", COMMISSION_RATE),
                                        ("Revenue target", TARGET)], 1):
        inputs.cell(r, 1, label)
        inputs.cell(r, 2, value)
    inputs["A1"].font = inputs["B1"].font = HEAD

    sales = wb.create_sheet("Sales")
    headers = ["Region", "Rep", "Units", "Unit Price", "Revenue", "Commission"]
    for c, h in enumerate(headers, 1):
        cell = sales.cell(1, c, h)
        cell.font, cell.fill = HEAD, FILL
    for i, (region, rep, units, price) in enumerate(DATA):
        r = FIRST + i
        sales.cell(r, 1, region)
        sales.cell(r, 2, rep)
        sales.cell(r, 3, units)
        sales.cell(r, 4, price)
        sales.cell(r, 5, f"=C{r}*D{r}")
        sales.cell(r, 6, f"=E{r}*CommissionRate")
        sales.cell(r, 5).number_format = sales.cell(r, 6).number_format = "#,##0.00"
    t = TOTAL_ROW
    sales.cell(t, 1, "Total").font = HEAD
    sales.cell(t, 3, f"=SUM(C{FIRST}:C{LAST})")
    sales.cell(t, 5, f"=SUM(E{FIRST}:E{LAST})")
    sales.cell(t, 6, f"=SUM(F{FIRST}:F{LAST})")
    for col in "ABCDEF":
        sales.column_dimensions[col].width = 14

    summary = wb.create_sheet("Summary")
    summary["A1"] = "Q3 Summary"
    summary["A1"].font = HEAD
    rows = [
        ("Total revenue", f"=Sales!E{t}"),
        ("Total commission", f"=Sales!E{t}" if buggy else f"=Sales!F{t}"),          # bug 1
        ("Average sale", "=AVERAGE(SalesRevenue)"),
        ("North revenue", f'=SUMIF(Sales!B{FIRST}:B{LAST},"North",Sales!E{FIRST}:E{LAST})'
         if buggy else f'=SUMIF(Sales!A{FIRST}:A{LAST},"North",Sales!E{FIRST}:E{LAST})'),  # bug 2
        ("Target met?", '=IF(B3>=Target,"Yes","No")'),
    ]
    for i, (label, formula) in enumerate(rows):
        summary.cell(3 + i, 1, label)
        summary.cell(3 + i, 2, formula)
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 14

    for name, ref in [("TaxRate", "Inputs!$B$2"), ("CommissionRate", "Inputs!$B$3"),
                      ("Target", "Inputs!$B$4"),
                      ("SalesRevenue", f"Sales!$E${FIRST}:$E${LAST}")]:
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    wb.save(path)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build(OUT / "sales_q3.xlsx")
    build(OUT / "sales_q3_buggy.xlsx", buggy=True)
    print(f"wrote {OUT / 'sales_q3.xlsx'} and {OUT / 'sales_q3_buggy.xlsx'}")
