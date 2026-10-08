import openpyxl, copy
wb = openpyxl.load_workbook('book.xlsx')
ws = wb['Sales']
ws.insert_rows(11)
for c in 'ABCDEF':
    ws[c + '11']._style = copy.copy(ws[c + '10']._style)
ws.row_dimensions[11].height = ws.row_dimensions[10].height
for c, v in zip('ABCD', ['West', 'Dana Lee', 40, 125]):
    ws[c + '11'] = v
ws['E11'] = '=C11*D11'
ws['F11'] = '=E11*CommissionRate'
for r in range(12, 23):
    ws[f'E{r}'] = f'=C{r}*D{r}'
    ws[f'F{r}'] = f'=E{r}*CommissionRate'
ws['C23'] = '=SUM(C2:C22)'
ws['E23'] = '=SUM(E2:E22)'
ws['F23'] = '=SUM(F2:F22)'
ws['E22'] = '=C22*D22'
wb.defined_names['SalesRevenue'].attr_text = 'Sales!$E$2:$E$22'
sm = wb['Summary']
sm['B3'] = '=Sales!E23'
sm['B4'] = '=Sales!F23'
sm['B6'] = '=SUMIF(Sales!A2:A22,"North",Sales!E2:E22)'
wb.save('book.xlsx')
for r in (10, 11, 12, 22, 23):
    print([ws.cell(r, c).value for c in range(1, 7)])
