#!/usr/bin/env python3
from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('/Users/shashankbhat/Projects/compiler_project/demo/runs/live-haiku/t2_insert_row__default__1/book.xlsx')

sales_sheet = wb['Sales']
print("All data rows in Sales sheet (showing row numbers and formulas):")
for row_idx in range(1, 24):
    cells = []
    for col_idx in range(1, 7):
        cell = sales_sheet.cell(row_idx, col_idx)
        if cell.value is not None:
            cells.append(str(cell.value))
        else:
            cells.append("")
    print(f"Row {row_idx:2d}: {cells}")

wb.close()
