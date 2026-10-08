#!/usr/bin/env python3
from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('/Users/shashankbhat/Projects/compiler_project/demo/runs/live-haiku/t2_insert_row__default__1/book.xlsx')

# List all named ranges
print("Named ranges in workbook:")
for name in wb.defined_names:
    defined_name = wb.defined_names[name]
    print(f"  Name: {name}")
    print(f"  Value: {defined_name.value}")
    print()

# Also check the Inputs sheet for parameters
inputs_sheet = wb['Inputs']
print("\n=== Inputs Sheet ===")
for row_idx in range(1, inputs_sheet.max_row + 1):
    row_data = []
    for col_idx in range(1, inputs_sheet.max_column + 1):
        cell = inputs_sheet.cell(row_idx, col_idx)
        if cell.value is not None:
            row_data.append(str(cell.value))
    if row_data:
        print(f"Row {row_idx}: {row_data}")

wb.close()
