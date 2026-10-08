#!/usr/bin/env python3
from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('/Users/shashankbhat/Projects/compiler_project/demo/runs/live-haiku/t2_insert_row__default__1/book.xlsx')

sales_sheet = wb['Sales']
print("Last 5 rows of Sales sheet:")
for row_idx in range(sales_sheet.max_row - 4, sales_sheet.max_row + 1):
    row_data = []
    for col_idx in range(1, sales_sheet.max_column + 1):
        cell = sales_sheet.cell(row_idx, col_idx)
        if cell.value is not None:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                row_data.append(f"[{cell.value}]")
            else:
                row_data.append(str(cell.value))
        else:
            row_data.append("")
    print(f"Row {row_idx}: {row_data}")

wb.close()
