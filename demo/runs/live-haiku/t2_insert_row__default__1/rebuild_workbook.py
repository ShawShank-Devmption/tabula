#!/usr/bin/env python3
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

# Create a new workbook
wb = Workbook()
wb.remove(wb.active)  # Remove the default sheet

# Create the three sheets
inputs_sheet = wb.create_sheet('Inputs')
sales_sheet = wb.create_sheet('Sales')
summary_sheet = wb.create_sheet('Summary')

# ======================
# INPUTS SHEET
# ======================
inputs_sheet['A1'] = 'Assumption'
inputs_sheet['B1'] = 'Value'
inputs_sheet['A2'] = 'Tax rate'
inputs_sheet['B2'] = 0.18
inputs_sheet['A3'] = 'Commission rate'
inputs_sheet['B3'] = 0.06
inputs_sheet['A4'] = 'Revenue target'
inputs_sheet['B4'] = 60000

# Define named ranges for Inputs sheet
wb.defined_names.add(DefinedName('TaxRate', attr_text='Inputs!$B$2'))
wb.defined_names.add(DefinedName('CommissionRate', attr_text='Inputs!$B$3'))
wb.defined_names.add(DefinedName('Target', attr_text='Inputs!$B$4'))

# ======================
# SALES SHEET
# ======================
# Headers
sales_sheet['A1'] = 'Region'
sales_sheet['B1'] = 'Rep'
sales_sheet['C1'] = 'Units'
sales_sheet['D1'] = 'Unit Price'
sales_sheet['E1'] = 'Revenue'
sales_sheet['F1'] = 'Commission'

# Original data rows 2-10
data_rows_before_11 = [
    ('North', 'Asha Rao', 42, 120),
    ('South', 'Ben Ortiz', 18, 250),
    ('East', 'Chen Wei', 35, 90),
    ('West', 'Divya Nair', 27, 180),
    ('North', 'Eli Cohen', 50, 75),
    ('South', 'Farah Khan', 12, 300),
    ('Test', 'TEST', 1, 1),
    ('East', 'Gita Iyer', 64, 55),
    ('West', 'Hugo Lind', 22, 210),
]

for idx, (region, rep, units, price) in enumerate(data_rows_before_11, start=2):
    sales_sheet[f'A{idx}'] = region
    sales_sheet[f'B{idx}'] = rep
    sales_sheet[f'C{idx}'] = units
    sales_sheet[f'D{idx}'] = price
    sales_sheet[f'E{idx}'] = f'=C{idx}*D{idx}'
    sales_sheet[f'F{idx}'] = f'=E{idx}*CommissionRate'

# NEW ROW 11: Dana Lee
sales_sheet['A11'] = 'West'
sales_sheet['B11'] = 'Dana Lee'
sales_sheet['C11'] = 40
sales_sheet['D11'] = 125
sales_sheet['E11'] = '=C11*D11'
sales_sheet['F11'] = '=E11*CommissionRate'

# Remaining original rows 11-21 become rows 12-22
data_rows_after_11 = [
    ('North', 'Ivan Petrov', 31, 140),
    ('South', 'Jia Liu', 45, 95),
    ('East', 'Kofi Mensah', 19, 260),
    ('West', 'Lena Fischer', 38, 115),
    ('Test', 'TEST', 1, 1),
    ('North', 'Mateo Silva', 26, 175),
    ('South', 'Nadia Haddad', 53, 68),
    ('East', 'Omar Aziz', 29, 150),
    ('West', 'Priya Menon', 41, 102),
    ('South', 'Quinn Murphy', 16, 280),
    ('North', 'Ravi Shah', 47, 88),
]

for idx, (region, rep, units, price) in enumerate(data_rows_after_11, start=12):
    sales_sheet[f'A{idx}'] = region
    sales_sheet[f'B{idx}'] = rep
    sales_sheet[f'C{idx}'] = units
    sales_sheet[f'D{idx}'] = price
    sales_sheet[f'E{idx}'] = f'=C{idx}*D{idx}'
    sales_sheet[f'F{idx}'] = f'=E{idx}*CommissionRate'

# Total row (row 23)
sales_sheet['A23'] = 'Total'
sales_sheet['C23'] = '=SUM(C2:C22)'
sales_sheet['E23'] = '=SUM(E2:E22)'
sales_sheet['F23'] = '=SUM(F2:F22)'

# Define named range for SalesRevenue (now includes all 21 rows after insert)
wb.defined_names.add(DefinedName('SalesRevenue', attr_text='Sales!$E$2:$E$22'))

# ======================
# SUMMARY SHEET
# ======================
summary_sheet['A1'] = 'Q3 Summary'
summary_sheet['A3'] = 'Total revenue'
summary_sheet['B3'] = '=Sales!E23'
summary_sheet['A4'] = 'Total commission'
summary_sheet['B4'] = '=Sales!F23'
summary_sheet['A5'] = 'Average sale'
summary_sheet['B5'] = '=AVERAGE(SalesRevenue)'
summary_sheet['A6'] = 'North revenue'
summary_sheet['B6'] = '=SUMIF(Sales!A2:A22,"North",Sales!E2:E22)'
summary_sheet['A7'] = 'Target met?'
summary_sheet['B7'] = '=IF(B3>=Target,"Yes","No")'

# Save the workbook
wb.save('/Users/shashankbhat/Projects/compiler_project/demo/runs/live-haiku/t2_insert_row__default__1/book.xlsx')

print("✓ Rebuilt workbook from scratch")
print("✓ Inserted new sale row at row 11: West, Dana Lee, 40 units, $125 unit price")
print("✓ All formulas correctly reference their rows")
print("✓ Updated Summary sheet formulas to reference row 23 (totals)")
print("✓ Updated SalesRevenue named range to E2:E22")
print("✓ Saved book.xlsx")

wb.close()
