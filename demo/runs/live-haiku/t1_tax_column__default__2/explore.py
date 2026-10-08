#!/usr/bin/env python3
import openpyxl

# Load the workbook
wb = openpyxl.load_workbook('book.xlsx')

# List all sheets
print("Sheets:", wb.sheetnames)

# Explore each sheet
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== {sheet_name} ===")
    print(f"Dimensions: {ws.dimensions}")

    # Show first 25 rows to understand structure
    for row in ws.iter_rows(min_row=1, max_row=25, values_only=False):
        row_data = []
        for cell in row:
            if cell.value is not None:
                row_data.append(f"{cell.coordinate}:{cell.value}")
        if row_data:
            print(row_data)
