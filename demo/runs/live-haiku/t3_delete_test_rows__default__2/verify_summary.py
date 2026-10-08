from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')

print("Summary sheet content:")
ws = wb['Summary']
for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
    values = [cell.value for cell in row]
    print(f"Row {row_idx}: {values}")

print("\n\nInputs sheet content:")
ws = wb['Inputs']
for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
    values = [cell.value for cell in row]
    print(f"Row {row_idx}: {values}")
