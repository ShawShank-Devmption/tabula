from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')
ws_summary = wb['Summary']

# Update Summary sheet formulas to reflect new row positions
# The Total row moved from row 22 to row 20
# Data range moved from A2:A21, E2:E21 to A2:A19, E2:E19

ws_summary['B3'].value = '=Sales!E20'  # Total revenue (was E22)
ws_summary['B4'].value = '=Sales!F20'  # Total commission (was F22)
ws_summary['B6'].value = '=SUMIF(Sales!A2:A19,"North",Sales!E2:E19)'  # North revenue

print("Summary sheet after updating references:")
for row_idx, row in enumerate(ws_summary.iter_rows(values_only=False), start=1):
    values = [cell.value for cell in row]
    print(f"Row {row_idx}: {values}")

# Save the workbook
wb.save('book.xlsx')
print("\nWorkbook saved successfully with updated Summary references!")
