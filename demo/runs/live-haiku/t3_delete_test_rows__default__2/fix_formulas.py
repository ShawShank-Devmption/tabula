from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')
ws = wb['Sales']

# Fix formulas for data rows (rows 2 through max_row - 1, since last row is Total)
max_data_row = ws.max_row - 1  # Exclude the Total row

for row_idx in range(2, max_data_row + 1):
    # Revenue: =C{row}*D{row}
    ws[f'E{row_idx}'].value = f'=C{row_idx}*D{row_idx}'
    # Commission: =E{row}*CommissionRate
    ws[f'F{row_idx}'].value = f'=E{row_idx}*CommissionRate'

# Fix the Total row (should be at row max_data_row + 1)
total_row = max_data_row + 1
ws[f'C{total_row}'].value = f'=SUM(C2:C{max_data_row})'
ws[f'E{total_row}'].value = f'=SUM(E2:E{max_data_row})'
ws[f'F{total_row}'].value = f'=SUM(F2:F{max_data_row})'

# Verify the result
print("Sales sheet after fixing formulas:")
for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
    values = [cell.value for cell in row]
    print(f"Row {row_idx}: {values}")

# Save the workbook
wb.save('book.xlsx')
print("\nWorkbook saved successfully with corrected formulas!")
