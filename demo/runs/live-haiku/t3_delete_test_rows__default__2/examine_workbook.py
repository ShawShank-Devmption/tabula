from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')
print("Sheet names:", wb.sheetnames)

# Get the Sales sheet
if 'Sales' in wb.sheetnames:
    ws = wb['Sales']
    print(f"\nSales sheet dimensions: {ws.dimensions}")
    print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

    # Print all rows to understand the structure
    print("\nSales sheet content:")
    for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
        values = [cell.value for cell in row]
        print(f"Row {row_idx}: {values}")
