from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')
ws = wb['Sales']

# Find and delete TEST rows (from bottom to top to avoid index shifting)
rows_to_delete = []
for row_idx in range(2, ws.max_row):  # Skip header, skip total row
    rep_cell = ws[f'B{row_idx}']
    if rep_cell.value == 'TEST':
        rows_to_delete.append(row_idx)

print(f"Found TEST rows at: {rows_to_delete}")

# Delete from bottom to top
for row_idx in sorted(rows_to_delete, reverse=True):
    print(f"Deleting row {row_idx}")
    ws.delete_rows(row_idx)

# Verify the result
print("\nSales sheet after deletion:")
for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
    values = [cell.value for cell in row]
    print(f"Row {row_idx}: {values}")

# Save the workbook
wb.save('book.xlsx')
print("\nWorkbook saved successfully!")
