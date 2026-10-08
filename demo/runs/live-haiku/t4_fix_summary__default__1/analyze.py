#!/usr/bin/env python3
from openpyxl import load_workbook

# Load with formulas
wb_formulas = load_workbook('book.xlsx')
ws_summary = wb_formulas['Summary']

print("SUMMARY SHEET - Current Formulas:")
print("-" * 80)
for row_num in range(3, 8):
    cell_a = ws_summary[f'A{row_num}'].value
    cell_b_value = ws_summary[f'B{row_num}'].value
    cell_b_type = ws_summary[f'B{row_num}'].data_type
    print(f"Row {row_num}: {cell_a}")
    print(f"  B{row_num}: {cell_b_value} (type: {cell_b_type})")

print("\n" + "="*80 + "\n")

# Load with data_only to see calculated values
wb_values = load_workbook('book.xlsx', data_only=True)
ws_sales = wb_values['Sales']
ws_summary_values = wb_values['Summary']

print("SALES SHEET - Data (rows 2-21):")
print("-" * 80)
print(f"{'Row':<5} {'Region':<10} {'Revenue':<15} {'Commission':<15}")
for row_num in range(2, 22):
    region = ws_sales[f'A{row_num}'].value
    revenue = ws_sales[f'E{row_num}'].value
    commission = ws_sales[f'F{row_num}'].value
    print(f"{row_num:<5} {str(region):<10} {revenue!s:<15} {commission!s:<15}")

print("\n" + "="*80 + "\n")
print("SUMMARY SHEET - Calculated Values:")
print("-" * 80)
for row_num in range(3, 8):
    label = ws_summary_values[f'A{row_num}'].value
    value = ws_summary_values[f'B{row_num}'].value
    print(f"{label}: {value}")

# Check what's in E22 in Sales sheet
print("\n" + "="*80 + "\n")
print("CHECKING E22 IN SALES:")
print(f"E22 value: {ws_sales['E22'].value}")
print(f"E22 data_type: {ws_sales['E22'].data_type}")

wb_formulas.close()
wb_values.close()
