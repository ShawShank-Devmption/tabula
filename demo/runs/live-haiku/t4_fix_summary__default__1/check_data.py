#!/usr/bin/env python3
from openpyxl import load_workbook

wb = load_workbook('book.xlsx', data_only=True)
ws_sales = wb['Sales']

print("SALES DATA (with calculated values):")
print("-" * 80)
print(f"{'Row':<5} {'Region':<10} {'Rep':<20} {'Units':<8} {'UnitPrice':<12} {'Revenue':<12} {'Commission':<12}")
print("-" * 80)

for row in ws_sales.iter_rows(min_row=2, max_row=21, min_col=1, max_col=6, values_only=True):
    print(f"{ws_sales._current_row:<5} {str(row[0]):<10} {str(row[1]):<20} {str(row[2]):<8} {str(row[3]):<12} {str(row[4]):<12} {str(row[5]):<12}")

print("\nSUMMARY SHEET VALUES:")
ws_summary = wb['Summary']
for row in ws_summary.iter_rows(min_row=3, max_row=7, min_col=1, max_col=2, values_only=True):
    print(f"{row[0]}: {row[1]}")

wb.close()
