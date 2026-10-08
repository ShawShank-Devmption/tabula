#!/usr/bin/env python3
from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')

# List all sheets
print("Sheets in workbook:", wb.sheetnames)
print("\n" + "="*80 + "\n")

# Examine the Summary sheet
if 'Summary' in wb.sheetnames:
    ws_summary = wb['Summary']
    print("SUMMARY SHEET:")
    print("-" * 80)
    for row in ws_summary.iter_rows(min_row=1, max_row=30, min_col=1, max_col=10, values_only=False):
        for cell in row:
            if cell.value is not None:
                print(f"{cell.coordinate}: {cell.value} (formula: {cell.data_type == 'f' and cell.value})")

print("\n" + "="*80 + "\n")

# Examine other sheets briefly
for sheet_name in wb.sheetnames:
    if sheet_name != 'Summary':
        print(f"\n{sheet_name.upper()} SHEET (first 20 rows):")
        print("-" * 80)
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=10, values_only=False):
            for cell in row:
                if cell.value is not None:
                    formula_str = f" (FORMULA: {cell.value})" if cell.data_type == 'f' else ""
                    print(f"{cell.coordinate}: {cell.value}{formula_str}")

wb.close()
