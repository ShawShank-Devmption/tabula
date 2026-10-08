#!/usr/bin/env python3
from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')

print("Current sheets:", wb.sheetnames)
print()

# Check if "Sales" sheet exists
if "Sales" not in wb.sheetnames:
    print("ERROR: 'Sales' sheet not found!")
    exit(1)

print("Defined names (named ranges) in workbook:")
try:
    for name_obj in wb.defined_names:
        print(f"  {name_obj.name}: {name_obj.value}")
except:
    print("  (No defined names found or error accessing them)")
print()

# Check for formulas referencing "Sales" sheet
print("Formulas referencing 'Sales' sheet:")
formula_count = 0
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                if 'Sales' in cell.value:
                    print(f"  {sheet}!{cell.coordinate}: {cell.value}")
                    formula_count += 1
if formula_count == 0:
    print("  No formulas found referencing 'Sales'")
print()

print("="*60)
print("Renaming 'Sales' to 'Q3 Sales'...")
print("="*60 + "\n")

# Rename the sheet
sales_sheet = wb["Sales"]
sales_sheet.title = "Q3 Sales"

print(f"Sheet renamed successfully!")
print(f"New sheets: {wb.sheetnames}")
print()

# Update defined names (named ranges) that reference the old sheet name
print("Updating defined names (named ranges)...")
updated_ranges = []
try:
    for name_obj in wb.defined_names:
        old_value = name_obj.value
        # Replace 'Sales!' with 'Q3 Sales!' in the reference
        new_value = old_value.replace("Sales!", "'Q3 Sales'!")
        if old_value != new_value:
            print(f"  {name_obj.name}: {old_value} -> {new_value}")
            name_obj.value = new_value
            updated_ranges.append(name_obj.name)
except:
    pass

if not updated_ranges:
    print("  No defined names needed updating")
print()

# Update formulas in all sheets
print("Updating formulas referencing 'Sales'...")
updated_formulas = []
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                if 'Sales!' in cell.value or 'Sales.' in cell.value:
                    old_formula = cell.value
                    # Replace Sales! with 'Q3 Sales'! (with quotes for the space)
                    new_formula = old_formula.replace("Sales!", "'Q3 Sales'!")
                    if old_formula != new_formula:
                        print(f"  {sheet}!{cell.coordinate}: Updated")
                        cell.value = new_formula
                        updated_formulas.append((sheet, cell.coordinate))

if not updated_formulas:
    print("  No formulas needed updating")
print()

# Save the workbook
print("Saving book.xlsx...")
wb.save('book.xlsx')
print("Done!")
