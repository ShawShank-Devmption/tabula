from openpyxl import load_workbook

wb = load_workbook('book.xlsx', data_only=False)

# Check the Sales sheet structure
if 'Sales' in wb.sheetnames:
    ws = wb['Sales']
    print("Sales sheet structure:")
    print("-" * 80)

    # Print headers and some rows to understand structure
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=25, min_col=1, max_col=6), 1):
        row_data = []
        for cell in row:
            if cell.value is not None:
                row_data.append(f"{cell.coordinate}: {cell.value}")
        if row_data:
            print(" | ".join(row_data))

    print("\n" + "=" * 80)
    print("\nRow 22 (totals row) details:")
    row_22 = ws[22]
    for cell in row_22:
        if cell.value is not None:
            if cell.data_type == 'f':
                print(f"{cell.coordinate}: {cell.value} (formula)")
            else:
                print(f"{cell.coordinate}: {cell.value}")

# Check Inputs sheet
print("\n" + "=" * 80)
if 'Inputs' in wb.sheetnames:
    ws = wb['Inputs']
    print("\nInputs sheet contents:")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            if cell.value is not None:
                if cell.data_type == 'f':
                    print(f"{cell.coordinate}: {cell.value} (formula)")
                else:
                    print(f"{cell.coordinate}: {cell.value}")
