from openpyxl import load_workbook

wb = load_workbook('book.xlsx')

print("Workbook-level defined names:")
if wb.defined_names:
    for name in wb.defined_names:
        print(f"  {name}: {wb.defined_names[name].value}")
else:
    print("  (none)")

print("\nSheet-level defined names:")
for ws_name in wb.sheetnames:
    ws = wb[ws_name]
    if ws.defined_names:
        print(f"  {ws_name}:")
        for name in ws.defined_names:
            print(f"    {name}: {ws.defined_names[name].value}")
    else:
        print(f"  {ws_name}: (none)")
