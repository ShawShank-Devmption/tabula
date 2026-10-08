from openpyxl import load_workbook

wb = load_workbook('book.xlsx')
ws = wb['Summary']

# Fix B4: Total commission should reference F22 (commission total), not E22 (revenue total)
ws['B4'].value = '=Sales!F22'

# Fix B6: North revenue should check column A (Region), not column B (Rep names)
ws['B6'].value = '=SUMIF(Sales!A2:A21,"North",Sales!E2:E21)'

print("Fixed formulas:")
print(f"B4 (Total commission): {ws['B4'].value}")
print(f"B6 (North revenue): {ws['B6'].value}")

# Save the workbook
wb.save('book.xlsx')
print("\nWorkbook saved successfully!")
