from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')

# List all sheet names
print("Sheets in workbook:", wb.sheetnames)
print()

# Examine the Summary sheet
if 'Summary' in wb.sheetnames:
    ws = wb['Summary']
    print("Summary sheet contents:")
    print("-" * 80)

    # Print all cells with data
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        row_data = []
        for cell in row:
            if cell.value is not None:
                # Show the cell reference, value, and formula if it has one
                if cell.data_type == 'f':
                    row_data.append(f"{cell.coordinate}: {cell.value} (formula)")
                else:
                    row_data.append(f"{cell.coordinate}: {cell.value}")
        if row_data:
            print(" | ".join(row_data))

    print()
    print("=" * 80)
    print("\nDetailed cell analysis:")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.value is not None:
                print(f"Cell {cell.coordinate}:")
                print(f"  Value: {cell.value}")
                print(f"  Type: {cell.data_type}")
                if cell.data_type == 'f':
                    print(f"  Formula: {cell.value}")
                print()
