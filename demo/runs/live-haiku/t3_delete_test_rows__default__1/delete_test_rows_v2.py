from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')
sales_ws = wb['Sales']
summary_ws = wb['Summary']

# Identify test rows and delete them
test_rows = []
for row_idx in range(2, sales_ws.max_row + 1):
    if sales_ws[f'B{row_idx}'].value == 'TEST':
        test_rows.append(row_idx)

print(f"Found test rows at: {test_rows}")

# Delete rows from bottom to top to avoid row number shifting
for row_idx in sorted(test_rows, reverse=True):
    sales_ws.delete_rows(row_idx, 1)

print(f"Deleted test rows. New max_row: {sales_ws.max_row}")

# Now rebuild all formulas in the data rows
# The data should now be in rows 2-19, with row 20 being totals
for row_idx in range(2, sales_ws.max_row):  # 2 to 19 (data rows)
    # Revenue column (E): Units * Unit Price
    sales_ws[f'E{row_idx}'].value = f'=C{row_idx}*D{row_idx}'
    # Commission column (F): Revenue * CommissionRate
    sales_ws[f'F{row_idx}'].value = f'=E{row_idx}*CommissionRate'

# Update the total row (now at row 20)
sales_ws['C20'].value = '=SUM(C2:C19)'
sales_ws['E20'].value = '=SUM(E2:E19)'
sales_ws['F20'].value = '=SUM(F2:F19)'

# Update the named range SalesRevenue from E2:E21 to E2:E19
for defined_name in wb.defined_names:
    if defined_name == 'SalesRevenue':
        wb.defined_names['SalesRevenue'].attr_text = 'Sales!$E$2:$E$19'

# Update formulas in Summary sheet
summary_ws['B3'].value = '=Sales!E20'  # Total revenue
summary_ws['B4'].value = '=Sales!F20'  # Total commission
summary_ws['B6'].value = '=SUMIF(Sales!A2:A19,"North",Sales!E2:E19)'  # North revenue

# Save the workbook
wb.save('book.xlsx')

print("Test rows deleted successfully!")
print("All formulas updated.")
