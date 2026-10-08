from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')

# Print sheet names
print("Sheet names:", wb.sheetnames)

# Get the Sales sheet
if 'Sales' in wb.sheetnames:
    ws = wb['Sales']
    print("\nSales sheet contents:")
    print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")

    # Print all rows to see the data
    for row_idx, row in enumerate(ws.iter_rows(values_only=False), 1):
        print(f"Row {row_idx}: ", end="")
        for cell in row:
            print(f"{cell.value} ", end="")
        print()
