from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')

# Get the Summary sheet
if 'Summary' in wb.sheetnames:
    ws = wb['Summary']
    print("Summary sheet contents:")
    print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")

    # Print all rows
    for row_idx, row in enumerate(ws.iter_rows(values_only=False), 1):
        print(f"Row {row_idx}: ", end="")
        for cell in row:
            print(f"{cell.value} ", end="")
        print()
