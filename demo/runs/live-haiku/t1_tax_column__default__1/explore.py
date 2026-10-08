#!/usr/bin/env python3
from openpyxl import load_workbook

# Load the workbook
wb = load_workbook('book.xlsx')

# Check available sheets
print("Sheets in workbook:", wb.sheetnames)
print()

# Explore Inputs sheet for tax rate
if 'Inputs' in wb.sheetnames:
    inputs_sheet = wb['Inputs']
    print("=== Inputs Sheet ===")
    for row in inputs_sheet.iter_rows(min_row=1, max_row=10, min_col=1, max_col=2, values_only=False):
        for cell in row:
            if cell.value is not None:
                print(f"{cell.coordinate}: {cell.value}")
    print()

# Explore Sales sheet
if 'Sales' in wb.sheetnames:
    sales_sheet = wb['Sales']
    print("=== Sales Sheet ===")
    print("First 5 rows:")
    for row in sales_sheet.iter_rows(min_row=1, max_row=5, min_col=1, max_col=10, values_only=False):
        row_data = []
        for cell in row:
            row_data.append(f"{cell.coordinate}:{cell.value}")
        print(" | ".join(row_data))
    print()
    print("Row 22 (Total row):")
    for row in sales_sheet.iter_rows(min_row=22, max_row=22, min_col=1, max_col=10, values_only=False):
        row_data = []
        for cell in row:
            row_data.append(f"{cell.coordinate}:{cell.value}")
        print(" | ".join(row_data))
    print()

# Explore Summary sheet
if 'Summary' in wb.sheetnames:
    summary_sheet = wb['Summary']
    print("=== Summary Sheet ===")
    for row in summary_sheet.iter_rows(min_row=1, max_row=10, min_col=1, max_col=2, values_only=False):
        for cell in row:
            if cell.value is not None:
                print(f"{cell.coordinate}: {cell.value}")
    print()
